from copy import deepcopy
from io import BytesIO
import unittest
from openpyxl import load_workbook
from post_evaluation.model_defaults import default_model_definition
from post_evaluation.result_state import describe_result
from post_evaluation.score_presentation import step_format
from post_evaluation.exports import make_result_tables, make_workbook
from post_evaluation.constants import PRODUCTS
from test_execution_boundary import calculate_sample, sample_records, rule_options


class ScorePresentationTests(unittest.TestCase):
    def definition(self):
        definition = default_model_definition()
        definition['dimensions'] = [d for d in definition['dimensions'] if d['key'] != 'ability']
        definition['dimensions'][0].update(weight=20, label='行业竞争表现')
        return definition

    def test_weighted_contribution_and_radar_use_the_same_full_score(self):
        definition = self.definition()
        indicator = definition['dimensions'][0]
        indicator['options'] = rule_options(indicator)
        indicator['options']['parsed_rule']['parameters'].update(first_score=20, rank_decrement=1)
        result = calculate_sample(definition)
        product = PRODUCTS[3]
        trace = next(t for t in result['traces'] if t['product'] == product and t['metric'] == 'market_share')
        row = trace['contribution']
        self.assertEqual(trace['indicator_label'], '行业竞争表现')
        self.assertEqual(trace['score'], 17)
        self.assertEqual(row['components'][0]['first_score'], 20)
        self.assertEqual(row['weighted_score'], 17)
        self.assertEqual(row['ratio'], .85)
        self.assertEqual(row, result['indicator_scores'][product][0])
        table = make_result_tables({'year':2025, 'model':{'definition':definition}}, result)[0]
        self.assertEqual(next(r for r in table['rows'] if r['product']==product)['values'][2], row['weighted_score'])
        for summary in result['model_summary']['rows']:
            self.assertAlmostEqual(sum(d['weighted_score'] for d in result['indicator_scores'][summary['product']]), summary['total'])

    def test_weight_changes_contribution_but_preserves_radar_ratio(self):
        definition = self.definition()
        before = calculate_sample(definition)
        definition['dimensions'][0]['weight'] = 30
        definition['dimensions'][1]['weight'] = 0
        after = calculate_sample(definition)
        for product in PRODUCTS:
            a, b = before['indicator_scores'][product][0], after['indicator_scores'][product][0]
            self.assertAlmostEqual(a['ratio'], b['ratio'])
            self.assertAlmostEqual(b['weighted_score'], a['weighted_score']*1.5)
            self.assertIsNone(after['indicator_scores'][product][1]['ratio'])

    def test_missing_and_zero_scores_are_distinct_and_history_is_not_mutated(self):
        definition = self.definition()
        records = [r for r in sample_records() if r['field'] != 'market_share:2025']
        partial = calculate_sample(definition, records)
        self.assertIsNone(partial['indicator_scores'][PRODUCTS[0]][0]['ratio'])
        historical = calculate_sample(definition)
        historical['scores'][PRODUCTS[0]]['market_share'] = 0
        historical['model_summary']['rows'][0]['dimensions'][0]['score'] = 0
        before = deepcopy(historical)
        decorated = describe_result(definition, historical)
        self.assertEqual(decorated['indicator_scores'][PRODUCTS[0]][0]['ratio'], 0)
        self.assertEqual(historical, before)
        self.assertEqual(decorated['model_summary'], before['model_summary'])

    def test_fraction_money_ratio_and_average_formats_match_excel_and_trace(self):
        definition = self.definition()
        records = sample_records()
        for record in records:
            if record['field'].startswith('market_share:'):
                record['value'] = .519 if record['product']==PRODUCTS[0] else .1
            if record['field'].startswith(('revenue:', 'research_')):
                record['unit'] = '万元'
        result = calculate_sample(definition, records)
        product = PRODUCTS[0]
        market = next(t for t in result['traces'] if t['product']==product and t['metric']=='market_share')
        revenue = next(t for t in result['traces'] if t['product']==product and t['metric']=='revenue')
        self.assertEqual(market['steps'][0]['value'], .519)
        self.assertEqual(market['steps'][0]['format'], 'percent')
        self.assertEqual(revenue['steps'][1]['unit'], '万元')
        self.assertEqual(revenue['steps'][2]['format'], 'percent')
        self.assertEqual(revenue['steps'][3]['format'], 'score')
        research = next(t for t in result['traces'] if t['product']==product and t['metric']=='self_research')
        self.assertEqual(research['steps'][9]['unit'], '万元')
        run = {'year':2025, 'model':{'definition':definition}}
        tables = make_result_tables(run, result)
        workbook = load_workbook(BytesIO(make_workbook(run,result)))
        self.assertEqual(workbook['重点产品-市场份额']['C3'].number_format, '0.00%')
        finance = next(t for t in tables if t['label']=='效益贡献')
        money_index = next(i for i,c in enumerate(finance['columns']) if c['label']=='2025年收入')
        self.assertEqual(finance['columns'][money_index]['unit'], '万元')
        self.assertIn('万元', workbook['重点产品-效益贡献'].cell(3,money_index+1).number_format)
        self.assertEqual(step_format('self_research',10,{}), {'format':'ratio','unit':'倍'})
        self.assertEqual(step_format('revenue',0,{})['unit'], '单位未标注')
        self.assertEqual(step_format('gross_margin',0,{})['format'], 'percent')

    def test_multi_metric_legacy_indicator_explains_average_of_ratios(self):
        definition = default_model_definition()
        indicator = definition['dimensions'][0]
        indicator.update(metric_keys=['market_share','quality'], weight=20)
        indicator.pop('calculation_type',None)
        raw = {'scores':{PRODUCTS[0]:{'market_share':8.5,'quality':7.5}}, 'traces':[],
               'model_summary':{'rows':[{'product':PRODUCTS[0],'dimensions':[{'key':indicator['key'],'score':16}]}]}}
        result = describe_result(definition,raw)
        row = result['indicator_scores'][PRODUCTS[0]][0]
        self.assertEqual([c['ratio'] for c in row['components']],[.85,.75])
        self.assertEqual(row['weighted_score'],16)
        self.assertEqual(row['ratio'],.8)
