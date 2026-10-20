from copy import deepcopy
from io import BytesIO
from threading import Event
import os
import tempfile
import time
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from post_evaluation import store, report_jobs
from post_evaluation.app import app
from post_evaluation.exports import make_report
from post_evaluation.report_analysis import report_evidence
from post_evaluation.report_insights import dimension_evidence, generate_report_draft, validate_section
from post_evaluation.report_preview import report_html
from post_evaluation.model_defaults import default_model_definition
from test_execution_boundary import calculate_sample


def section(facts):
    return {"observations": [{"title": "产品间表现存在差异", "text": "根据本批次数据，需结合原值和变化趋势评价产品表现。",
                              "evidence_ids": [facts[0]["id"]]}],
            "recommendation": "建议优先核实相对薄弱指标，并跟踪改善效果。", "evidence_ids": [facts[0]["id"]]}


def full_draft(run, result, progress=lambda *args: None):
    evidence = report_evidence(run, result)
    scopes = dimension_evidence(run, result, evidence)
    return {"version": 2, "signature": evidence["signature"], "provider": "test", "model": "test", "generated_at": "2026-10-20T00:00:00Z",
            "overview": section(scopes[0]["facts"]),
            "dimensions": [{"key": scope["key"], "label": scope["label"], **section(scope["facts"])} for scope in scopes if scope["facts"]],
            "products": [{"product": row["product"], "evaluation": "产品各项指标表现存在差异。", "recommendation": "建议核实薄弱指标，跟踪变化。", "evidence_ids": [row["facts"][0]["id"]]} for row in evidence["products"]]}


class ReportWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.definition = default_model_definition()
        cls.result = calculate_sample(cls.definition)
        cls.report_run = {"year": 2025, "data_year": 2025, "model": {"definition": cls.definition}, "notes": {}, "files": []}

    def test_dimension_evidence_is_full_cohort_and_excludes_simulated(self):
        evidence = report_evidence(self.report_run, self.result)
        scopes = dimension_evidence(self.report_run, self.result, evidence)
        self.assertEqual(len(scopes), len(self.definition["categories"]))
        for scope in scopes:
            self.assertEqual(len({f["product"] for f in scope["facts"]}), 13)
            self.assertFalse(any(f["metric"] in self.result["simulated_metrics"] for f in scope["facts"]))
            for stat in scope["statistics"]:
                self.assertEqual(stat["sample_count"], stat["positive_count"] + stat["negative_count"] + stat["zero_count"])
                self.assertEqual(stat["sample_count"], 13)
        self.assertTrue(any(s["year_comparisons"] for s in scopes))
        for scope in scopes:
            if scope["label"] in {"产品三效", "品质服务", "自研能力"}:
                self.assertTrue(all("available_sample_total" not in s for s in scope["statistics"]))

    def test_dimension_and_overview_generation_use_facts_not_product_prose(self):
        progress = []
        def generate(context, facts):
            self.assertNotIn("evaluation", context)
            return section(facts)
        draft = full_draft(self.report_run, self.result)
        with patch("post_evaluation.report_insights._generate_section", side_effect=generate), patch("post_evaluation.report_insights.generate_evaluation_draft", return_value=draft):
            output = generate_report_draft(self.report_run, self.result, lambda *args: progress.append(args))
        self.assertEqual(output["version"], 2)
        self.assertEqual(len(output["dimensions"]), 5)
        self.assertIn(("summary", 0, 1), progress)
        facts = report_evidence(self.report_run, self.result)["products"][0]["facts"]
        broken = section(facts)
        broken["observations"][0]["evidence_ids"] = ["other-dimension"]
        with self.assertRaises(ValueError): validate_section(broken, facts)

    def test_report_outline_preview_and_download_snapshot(self):
        draft = full_draft(self.report_run, self.result)
        content = make_report({**self.report_run, "notes": {"evaluation_draft": draft}}, self.result)
        document = Document(BytesIO(content))
        outline = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(outline[:3], ["一 总体结论", "二 分维度分析", "三 分产品评价和提升建议"])
        self.assertEqual(len(document.inline_shapes), 14)
        preview = report_html(content)
        self.assertIn("三 分产品评价和提升建议", preview)
        self.assertEqual(preview.count('<img alt="报告图表"'), 14)
        self.assertNotIn("textarea", preview)
        self.assertNotRegex(preview, r"P\d+D\d+C\d+")
        doc = Document()
        doc.add_paragraph('<script>alert("x")</script>')
        raw = BytesIO();doc.save(raw)
        html = report_html(raw.getvalue())
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)

    def wait_done(self, client, path):
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            state = client.get(path + "/ai-report").json()
            if state["status"] not in {"queued", "running"}: return state
            time.sleep(.02)
        self.fail("background generation did not finish")

    def test_jobs_are_idempotent_keep_old_snapshot_on_failure_and_guard_stale_data(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, PRODUCT_EVAL_DATA_DIR=directory):
            client = TestClient(app)
            run = store.create_run(2025, "后台生成验证")
            path = f"/api/runs/{run['id']}"
            self.assertEqual(client.post(path + "/ai-report").status_code, 409)
            store.update_run_json(run["id"], "result", self.result)
            run = store.get_run(run["id"])
            store.save_report(run, run["notes"], b"old report")
            release = Event()
            def slow(run, result, progress):
                progress("dimensions", 1, 5)
                release.wait(5)
                return full_draft(run, result)
            with patch("post_evaluation.report_jobs.llm_client.configured", return_value=True), patch("post_evaluation.report_jobs.generate_report_draft", side_effect=slow) as generator, patch("post_evaluation.report_jobs.make_report", return_value=b"new report"):
                self.assertEqual(client.post(path + "/ai-report").status_code, 202)
                self.assertEqual(client.post(path + "/ai-report").status_code, 202)
                self.assertEqual(client.get(path + "/report.docx").content, b"old report")
                release.set()
                self.assertEqual(self.wait_done(client, path)["status"], "completed")
                self.assertEqual(generator.call_count, 1)
                downloaded = client.get(path + "/report.docx")
                self.assertEqual(downloaded.content, b"new report")
                self.assertEqual(int(downloaded.headers["content-length"]), len(b"new report"))
            with patch("post_evaluation.report_jobs.llm_client.configured", return_value=True), patch("post_evaluation.report_jobs.generate_report_draft", side_effect=RuntimeError("private error")):
                client.post(path + "/ai-report")
                state = self.wait_done(client, path)
                self.assertEqual(state["status"], "failed")
                self.assertNotIn("private error", state["message"])
                self.assertEqual(client.get(path + "/report.docx").content, b"new report")
            stale = store.get_run(run["id"])
            store.update_run_json(run["id"], "notes", {"analysis": "新说明"})
            with self.assertRaisesRegex(ValueError, "已变化"):
                store.save_report(stale, {}, b"stale", expected_fingerprint=store.report_fingerprint(stale))
            self.assertEqual(store.report_content(run["id"]), b"new report")

    def test_preview_requires_current_report_and_never_calls_ai(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, PRODUCT_EVAL_DATA_DIR=directory):
            client = TestClient(app)
            run = store.create_run(2025, "预览快照")
            path = f"/api/runs/{run['id']}"
            self.assertEqual(client.get(path + "/report-preview").status_code, 409)
            store.update_run_json(run["id"], "result", self.result)
            run = store.get_run(run["id"])
            doc = Document();doc.add_paragraph("保存的报告内容")
            content = BytesIO();doc.save(content)
            store.save_report(run, run["notes"], content.getvalue())
            with patch("post_evaluation.report_jobs.generate_report_draft") as ai:
                response = client.get(path + "/report-preview")
                self.assertEqual(response.status_code, 200)
                self.assertIn("保存的报告内容", response.text)
                self.assertEqual(client.get(path + "/ai-report").json()["status"], "completed")
                ai.assert_not_called()
            store.update_run_json(run["id"], "result", {"changed": True})
            self.assertEqual(client.get(path + "/report-preview").status_code, 409)
