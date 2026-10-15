from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from post_evaluation import store
from post_evaluation.app import app
from post_evaluation.calculation_types import bindings, calculation_types, local_rule, raw_fields
from post_evaluation.constants import ECONOMIC_FIELDS, PRODUCTS, SELF_RESEARCH_COST_FIELDS
from post_evaluation.exports import make_report, make_workbook
from post_evaluation.model_defaults import default_model_definition
from post_evaluation.model_engine import simulate, validate_definition
from post_evaluation.scoring import calculate


def rule_options(indicator):
    type_id = bindings(indicator)[0][1]
    return {"parsed_rule": local_rule(type_id, indicator["rule_summary"]),
            "rule_parser": {"source_summary": indicator["rule_summary"]}}


def sample_records(year=2025):
    """Independent raw inputs: product 1 beats product 13 for every real algorithm."""
    values = {}
    for type_id in [item["id"] for item in calculation_types()]:
        for field in raw_fields(type_id, year):
            key, period = field.split(":")
            for index, product in enumerate(PRODUCTS):
                if key in ECONOMIC_FIELDS:
                    value = 10 if int(period) == year - 1 else 40 - index
                elif key == "self_owned_ratio":
                    value = .9 - .03 * index
                elif key in SELF_RESEARCH_COST_FIELDS:
                    value = 1
                elif key == "research_revenue":
                    value = 100 - index
                else:
                    value = 20 - index
                values[(field, product)] = {"field": field, "product": product, "value": value,
                                          "file_id": "sample", "cell": f"C{index+2}", "source_name": "独立样例.xlsx"}
    return list(values.values())


def calculate_sample(definition, records=None):
    file = {"id": "sample", "path": "unused", "role": "generic", "mapping": {},
            "targets": [item["key"] for item in definition["dimensions"]]}
    with patch("post_evaluation.scoring.extract_records", return_value=(records if records is not None else sample_records(), [])):
        return calculate(2025, [file], {}, definition)


class ExecutionBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, PRODUCT_EVAL_DATA_DIR=self.directory.name, GEMINI_API_KEY="")
        self.environment.start()
        self.client = TestClient(app)
        self.original = self.client.get("/api/models").json()[0]
        self.draft = self.client.post("/api/models", json={"name": "边界验证", "version": "CHECK-DRAFT",
                                                          "clone_from": self.original["id"]}).json()

    def tearDown(self):
        self.environment.stop()
        self.directory.cleanup()

    def publish(self, definition):
        return self.client.post(f"/api/models/{self.draft['id']}/publish", json={"definition": definition})

    def test_empty_unknown_and_malformed_calculators_are_rejected_without_writes(self):
        cases = [[], ["unknown"], [None]]
        for metrics in cases:
            with self.subTest(metrics=metrics):
                definition = default_model_definition()
                definition["dimensions"][0]["metric_keys"] = metrics
                definition["dimensions"][0].pop("calculation_type", None)
                before = store.get_model(self.draft["id"])
                response = self.publish(definition)
                self.assertEqual(response.status_code, 400, response.text)
                self.assertIn("计算", response.json()["detail"])
                self.assertEqual(store.get_model(self.draft["id"]), before)
                with self.assertRaises(ValueError):
                    simulate(definition, {})
        for field, value in (("calculation_type", "invented"), ("options", {"algorithm": "invented"})):
            definition = default_model_definition()
            definition["dimensions"][0][field] = value
            self.assertEqual(self.publish(definition).status_code, 400)
        for value in (None, "invalid"):
            definition = default_model_definition()
            definition["dimensions"][0] = value
            self.assertEqual(self.publish(definition).status_code, 400)

    def test_invalid_numbers_and_parameters_are_not_executable(self):
        for value in (True, False, float("nan"), float("inf"), 10**1000, -1):
            with self.subTest(value=value):
                definition = default_model_definition()
                definition["ranking"]["first_score"] = value
                self.assertTrue(validate_definition(definition))
                definition = default_model_definition()
                definition["dimensions"][0]["weight"] = value
                self.assertTrue(validate_definition(definition))
        for params in ({"first_score": 0}, {"rank_decrement": -1}, {"cap_percent": 120},
                       {"growth_denominator_uses_absolute_value": "yes"}, {"unknown": 3}):
            definition = default_model_definition()
            indicator = definition["dimensions"][0]
            indicator["options"] = rule_options(indicator)
            indicator["options"]["parsed_rule"]["parameters"].update(params)
            response = self.publish(definition)
            self.assertEqual(response.status_code, 400, response.text)
        for value in ("invalid", [1], False):
            definition = default_model_definition()
            definition["ranking"] = value
            response = self.client.post("/api/models/validate", json={"definition": definition})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertTrue(response.json()["errors"])

    def test_legacy_invalid_draft_simulation_returns_business_error(self):
        run = store.create_run(2025, "旧草稿模拟边界")
        definition = default_model_definition()
        definition["dimensions"][0]["metric_keys"] = []
        with store.connect() as db:
            db.execute("UPDATE model_versions SET definition=? WHERE id=?", (json.dumps(definition), self.draft["id"]))
        response = self.client.post(f"/api/models/{self.draft['id']}/simulate", json={"run_id": run["id"]})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("计算标识", response.json()["detail"])

    def test_unsupported_steps_order_ties_window_and_stale_summary_block_publication(self):
        for mutate in (
            lambda r: r["steps"].append({"order": 4, "operation": "custom", "description": "人工加分", "formula": "value*2"}),
            lambda r: r["steps"].reverse(),
            lambda r: r.update(tie_method="dense"),
            lambda r: r.update(time_window="custom"),
            lambda r: r.update(ranking_direction="unclear"),
            lambda r: r.update(warnings=["分母尚不明确"]),
            lambda r: r["steps"][2].update(formula="value*2"),
            lambda r: r.update(normalized_summary="所有产品都直接得满分"),
        ):
            definition = default_model_definition()
            indicator = definition["dimensions"][0]
            indicator["options"] = rule_options(indicator)
            mutate(indicator["options"]["parsed_rule"])
            self.assertEqual(self.publish(definition).status_code, 400)
        definition = default_model_definition()
        indicator = definition["dimensions"][0]
        indicator["options"] = rule_options(indicator)
        indicator["rule_summary"] = indicator["rule_summary"].replace("降序", "升序")
        self.assertEqual(self.publish(definition).status_code, 400)
        indicator["options"] = {}
        indicator["rule_summary"] = "取客户投诉次数，低于5次加2分。"
        self.assertEqual(self.publish(definition).status_code, 400)

    def test_local_templates_publish_offline_and_keep_original_model(self):
        definition = default_model_definition()
        definition["dimensions"][0]["weight"] = 20
        definition["dimensions"][1]["weight"] = 0
        with patch("post_evaluation.app.parse_rule_summary", side_effect=AssertionError("不应联网")):
            for indicator in definition["dimensions"]:
                response = self.client.post("/api/rules/parse", json={"dimension_label": indicator["label"],
                                            "calculation_type": bindings(indicator)[0][1], "summary": indicator["rule_summary"]})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["provider"], "local")
            self.assertEqual(self.publish(definition).status_code, 200)
        self.assertEqual(store.get_model(self.original["id"]), self.original)

    def test_explicit_parameters_survive_subsequent_model_reads(self):
        definition = default_model_definition()
        economic = definition["dimensions"][1]
        economic["rule_summary"] = economic["rule_summary"].replace("ABS（上年值）", "（上年值）")
        economic["options"]["growth_denominator_absolute"] = False
        response = self.publish(definition)
        self.assertEqual(response.status_code, 200, response.text)
        first = store.get_model(self.draft["id"])
        second = store.get_model(self.draft["id"])
        self.assertFalse(second["definition"]["dimensions"][1]["options"]["growth_denominator_absolute"])
        self.assertEqual(first, second)

    def test_new_instance_requires_supported_type_and_uses_its_raw_source(self):
        definition = default_model_definition()
        indicator = definition["dimensions"][0]
        indicator.update(key="new_market", label="低份额关注度", calculation_type="market_share", metric_keys=["new_output"],
                         rule_summary="取最新年度市场份额，升序排名赋分。")
        self.assertEqual(self.publish(definition).status_code, 200)
        from post_evaluation.constants import input_specs_for_dimension
        self.assertEqual(input_specs_for_dimension(indicator, 2025), [{"key": "market_share:2025", "label": "2025年市场份额", "required": True}])
        result = calculate_sample(definition)
        self.assertEqual(result["scores"][PRODUCTS[-1]]["new_output"], 10)
        self.assertNotIn("market_share", result["scores"][PRODUCTS[-1]])
        trace = next(t for t in result["traces"] if t["metric"] == "new_output")
        self.assertEqual(trace["steps"][0]["source"]["source_name"], "独立样例.xlsx")

    def test_all_supported_real_types_keep_default_formula_and_trace(self):
        result = calculate_sample(default_model_definition())
        for type_id in [item["id"] for item in calculation_types() if item["id"] != "ability"]:
            with self.subTest(type_id=type_id):
                self.assertEqual(result["scores"][PRODUCTS[0]][type_id], 10)
                self.assertEqual(result["scores"][PRODUCTS[-1]][type_id], 4)
                trace = next(t for t in result["traces"] if t["metric"] == type_id and t["product"] == PRODUCTS[0])
                self.assertTrue(any(step.get("source") for step in trace["steps"]))
        self.assertEqual(result["result_status"], "trial")

    def test_self_research_workbook_preserves_card_and_onepark_exceptions(self):
        definition = default_model_definition()
        indicator = next(item for item in definition["dimensions"] if item["key"] == "self_research")
        indicator["weight"] = 100
        definition["dimensions"] = [indicator]
        fields = raw_fields("self_research", 2025)
        wb = Workbook(); ws = wb.active
        ws.append(["产品线", *fields])
        for product in PRODUCTS:
            name = "物联网卡结算" if product == "物联卡" else product
            ratio = .146 if product == "OnePark" else 1 if product == "物联卡" else .5
            ws.append([name, ratio, 100, *([1] * 7)])
        ws.append(["OnePark（结算）", .99, 100, *([1] * 7)])
        ws.append(["芯片（结算）", .99, 100, *([1] * 7)])
        path = Path(self.directory.name) / "例外口径原表.xlsx"
        wb.save(path)
        from openpyxl.utils import get_column_letter
        file = {"id": "exceptions", "path": str(path), "role": "self_research", "targets": ["self_research"],
                "mapping": {"sheet": ws.title, "header_row": 1, "product_col": "A", "fields": {field: get_column_letter(index+2) for index, field in enumerate(fields)}}}
        result = calculate(2025, [file], {}, definition)
        self.assertFalse(result["issues"])
        traces = {t["product"]: t for t in result["traces"]}
        self.assertEqual(traces["OnePark"]["steps"][0]["value"], .146)
        self.assertEqual(traces["OnePark"]["steps"][0]["source"]["source_name"], "OnePark")
        self.assertEqual(traces["物联卡"]["steps"][0]["value"], 1)
        self.assertEqual(traces["物联卡"]["steps"][0]["source"]["source_name"], "物联网卡结算")
        self.assertEqual(traces["物联卡"]["steps"][10]["value"], 100 / 7)

    def test_direction_changes_real_scores_and_competition_ties(self):
        definition = default_model_definition()
        descending = calculate_sample(definition)
        indicator = definition["dimensions"][0]
        indicator["rule_summary"] = indicator["rule_summary"].replace("降序", "升序")
        ascending = calculate_sample(definition)
        self.assertEqual(descending["scores"][PRODUCTS[0]]["market_share"], 10)
        self.assertEqual(ascending["scores"][PRODUCTS[0]]["market_share"], 4)
        self.assertEqual(ascending["scores"][PRODUCTS[-1]]["market_share"], 10)
        records = sample_records()
        for record in records:
            if record["field"] == "market_share:2025" and record["product"] in PRODUCTS[-2:]:
                record["value"] = 1
        ties = calculate_sample(definition, records)
        ranks = {t["product"]: t["rank"] for t in ties["traces"] if t["metric"] == "market_share"}
        self.assertEqual([ranks[p] for p in PRODUCTS[-3:]], [3, 1, 1])

    def test_parsed_score_growth_and_cap_parameters_reach_execution(self):
        definition = default_model_definition()
        market, economic, research = definition["dimensions"][0], definition["dimensions"][1], definition["dimensions"][7]
        economic["rule_summary"] = economic["rule_summary"].replace("ABS（上年值）", "（上年值）")
        for indicator in (market, economic, research):
            indicator["options"] = rule_options(indicator)
        market["options"]["parsed_rule"]["parameters"].update(first_score=20, rank_decrement=1)
        economic["options"]["parsed_rule"]["parameters"]["growth_denominator_uses_absolute_value"] = False
        research["options"]["parsed_rule"]["parameters"]["cap_percent"] = 50
        records = sample_records()
        for record in records:
            if record["field"] == "revenue:2024":
                record["value"] = -10
        result = calculate_sample(definition, records)
        self.assertEqual(result["scores"][PRODUCTS[0]]["market_share"], 20)
        contribution = next(item for item in result["model_summary"]["rows"][0]["dimensions"] if item["key"] == "market")
        self.assertEqual(contribution["score"], 10)
        trace = next(t for t in result["traces"] if t["metric"] == "revenue" and t["product"] == PRODUCTS[0])
        self.assertEqual(trace["steps"][2]["value"], -5)
        trace = next(t for t in result["traces"] if t["metric"] == "self_research" and t["product"] == PRODUCTS[0])
        self.assertEqual(trace["steps"][0]["value"], .5)
        self.assertEqual(trace["steps"][10]["value"], 100 / 7)

    def test_dynamic_coverage_nine_indicators_weight_missing_and_real_status(self):
        definition = default_model_definition()
        definition["dimensions"] = [item for item in definition["dimensions"] if item["key"] != "ability"]
        definition["dimensions"][0]["weight"] = 20
        result = calculate_sample(definition)
        self.assertEqual((result["indicator_count"], result["calculated_indicator_count"], result["covered_points"]), (9, 9, 100))
        self.assertEqual(result["simulated_metrics"], [])
        self.assertTrue(result["official_total_available"])
        self.assertEqual(result["result_status"], "complete")
        records = [r for r in sample_records() if r["field"] != "market_share:2025"]
        partial = calculate_sample(definition, records)
        self.assertEqual(partial["covered_points"], 80)
        self.assertEqual(partial["calculated_indicator_count"], 8)
        self.assertEqual(partial["missing_indicators"], ["市场份额"])
        self.assertEqual(partial["result_status"], "partial")
        self.assertFalse(partial["official_total_available"])
        self.assertTrue(all(row["total"] is None for row in partial["model_summary"]["rows"]))

    def test_workbook_report_share_coverage_and_simulation_labels(self):
        for remove_mock, remove_source in ((False, False), (False, True), (True, False), (True, True)):
            with self.subTest(remove_mock=remove_mock, remove_source=remove_source):
                definition = default_model_definition()
                if remove_mock:
                    definition["dimensions"] = [item for item in definition["dimensions"] if item["key"] != "ability"]
                    definition["dimensions"][0]["weight"] = 20
                records = sample_records()
                if remove_source:
                    records = [r for r in records if r["field"] != "market_share:2025"]
                result = calculate_sample(definition, records)
                run = {"year": 2025, "name": "输出核对", "files": [], "notes": {}, "model": {"definition": definition}}
                wb = load_workbook(BytesIO(make_workbook(run, result)))
                cells = [str(cell.value) for row in wb["重点产品-雷达图维度"] for cell in row if cell.value is not None]
                self.assertIn(result["status_note"], cells)
                document = Document(BytesIO(make_report(run, result)))
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn(result["status_note"], text)
                self.assertIn(f"{result['calculated_indicator_count']}/{result['indicator_count']} 项指标", text)
                self.assertEqual(len(document.tables[0].columns), 1 + result["calculated_indicator_count"] + (0 if remove_source else 2))
                if remove_mock:
                    self.assertNotIn("能力评测为模拟值", text)

    def test_reading_legacy_coverage_does_not_change_stored_result_or_report(self):
        definition = default_model_definition()
        result = calculate_sample(definition)
        run = store.create_run(2025, "历史结果")
        old_result = {key: deepcopy(value) for key, value in result.items() if key not in {"indicators", "indicator_count", "calculated_indicator_count", "result_status", "status_note", "weight_total"}}
        old_result["covered_points"] = 999
        store.update_run_json(run["id"], "result", old_result)
        response = self.client.get(f"/api/runs/{run['id']}")
        self.assertEqual(response.json()["result"]["covered_points"], 100)
        self.assertEqual(store.get_run(run["id"])["result"], old_result)
        with patch("post_evaluation.app.make_report", return_value=b"snapshot"):
            generated = self.client.post(f"/api/runs/{run['id']}/report", json={})
        self.assertIsNotNone(generated.json()["report"])
        self.assertEqual(self.client.get(f"/api/runs/{run['id']}/report.docx").content, b"snapshot")

    def test_real_workbook_flow_binds_new_instance_and_calculates_ascending(self):
        definition = default_model_definition()
        indicator = definition["dimensions"][0]
        indicator.update(key="ascending_market", label="份额观察", metric_keys=["ascending_output"], calculation_type="market_share",
                         weight=100, rule_summary="取最新年度市场份额，升序排名赋分。")
        definition["dimensions"] = [indicator]
        published = self.publish(definition)
        self.assertEqual(published.status_code, 200, published.text)
        run = self.client.post("/api/runs", json={"year": 2025, "name": "原表升序核对", "model_id": self.draft["id"]}).json()
        wb = Workbook(); ws = wb.active
        ws.append(["公司产品", "2025年市场份额"])
        for index, product in enumerate(PRODUCTS):
            ws.append([product, (index + 1) / 100])
        content = BytesIO(); wb.save(content)
        uploaded = self.client.post(f"/api/runs/{run['id']}/files", files={"files": ("市场份额.xlsx", content.getvalue())})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        self.assertEqual(uploaded.json()["files"][0]["targets"], ["ascending_market"])
        validated = self.client.post(f"/api/runs/{run['id']}/validate")
        self.assertEqual(validated.json()["validation"]["status"], "passed", validated.text)
        calculated = self.client.post(f"/api/runs/{run['id']}/calculate")
        self.assertEqual(calculated.status_code, 200, calculated.text)
        result = calculated.json()
        self.assertEqual(result["scores"][PRODUCTS[0]]["ascending_output"], 10)
        self.assertEqual(result["model_summary"]["rows"][0]["total"], 100)
        self.assertTrue(result["official_total_available"])


if __name__ == "__main__":
    unittest.main()
