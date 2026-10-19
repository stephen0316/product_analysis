import unittest
from io import BytesIO

from openpyxl import Workbook

from post_evaluation.ingest import _find_finance_mapping, _find_year_column, auto_mapping, detect_metrics, extract_records


class QualityMappingTests(unittest.TestCase):
    def source(self, headers, *, title="重点产品-品质服务", sheet_name="Sheet1", second_sheet=False):
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        sheet.append([title])
        sheet.append(headers)
        sheet.append([1, "物联卡", 96.7, 94.3, 90.0][:len(headers)])
        if second_sheet:
            other = book.copy_worksheet(sheet)
            other.title = "第二个来源"
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        return buffer.getvalue()

    def test_raw_satisfaction_header_detects_maps_and_traces(self):
        content = self.source(["序号", "产品线", "产品推广满意度"])
        self.assertEqual(detect_metrics(BytesIO(content)), ["quality"])
        mapping = auto_mapping(BytesIO(content), ["quality"], 2025)
        self.assertEqual(mapping, {"sheet": "Sheet1", "header_row": 2, "product_col": "B", "fields": {"quality:2025": "C"}})
        records, issues = extract_records(BytesIO(content), "source", "generic", mapping)
        self.assertEqual(issues, [])
        self.assertEqual(records[0]["product"], "物联卡")
        self.assertEqual(records[0]["value"], 96.7)
        self.assertEqual(records[0]["cell"], "Sheet1!C3")

    def test_header_whitespace_and_years_follow_assessment_year(self):
        content = self.source(["序号", "产 品\n线", "2025年产品推广满意度", "产品推广\n满意度（2026年）"])
        self.assertEqual(detect_metrics(BytesIO(content)), ["quality"])
        mapping = auto_mapping(BytesIO(content), ["quality"], 2026)
        self.assertEqual(mapping["fields"], {"quality:2026": "D"})
        self.assertIsNone(auto_mapping(BytesIO(content), ["quality"], 2027)["fields"]["quality:2027"])

    def test_undated_header_uses_assessment_year_without_calendar_guess(self):
        content = self.source(["序号", "产品线", "产品推广满意度"])
        self.assertEqual(auto_mapping(BytesIO(content), ["quality"], 2027)["fields"], {"quality:2027": "C"})

    def test_wrong_year_in_title_or_sheet_is_not_silently_relabelled(self):
        for kwargs in ({"title": "2025年重点产品-品质服务"}, {"sheet_name": "2025年品质服务"}):
            with self.subTest(**kwargs):
                content = self.source(["序号", "产品线", "产品推广满意度"], **kwargs)
                self.assertIsNone(auto_mapping(BytesIO(content), ["quality"], 2026)["fields"]["quality:2026"])

    def test_ambiguous_columns_or_sheets_require_confirmation(self):
        cases = [self.source(["序号", "产品线", "产品推广满意度", "产品推广满意度"]),
                 self.source(["序号", "产品线", "产品推广满意度"], second_sheet=True)]
        for content in cases:
            self.assertEqual(detect_metrics(BytesIO(content)), ["quality"])
            self.assertIsNone(auto_mapping(BytesIO(content), ["quality"], 2025)["fields"]["quality:2025"])

    def test_title_or_derived_score_alone_is_not_a_raw_input(self):
        for header in ("排名", "产品推广满意度得分", "产品推广满意度排名", "客户满意度"):
            with self.subTest(header=header):
                self.assertEqual(detect_metrics(BytesIO(self.source(["序号", "产品线", header]))), [])


class DynamicYearMappingTests(unittest.TestCase):
    def test_finance_columns_follow_source_year_headers(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([])
        sheet.append([
            "产品线", "产品收入", None, "产品毛利", None,
            "产品毛利率", None, "营业利润", None, None, None,
        ])
        sheet.append([
            None, "25年", "26年", "25年", "26年", "25年", "26年",
            "25年营业利润额", "26年营业利润额", "25年营业利润率", "26年营业利润率",
        ])

        mapping = _find_finance_mapping(sheet, 2026)

        self.assertEqual(mapping["header_row"], 3)
        self.assertEqual(mapping["product_col"], "A")
        self.assertEqual(mapping["fields"], {
            "revenue:2025": "B", "revenue:2026": "C",
            "gross_profit:2025": "D", "gross_profit:2026": "E",
            "gross_margin:2025": "F", "gross_margin:2026": "G",
            "operating_profit:2025": "H", "operating_profit:2026": "I",
            "operating_margin:2025": "J", "operating_margin:2026": "K",
        })

    def test_renamed_share_header_can_be_found_by_year(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["公司产品", "行业份额（2025年）", "行业份额（2026年）"])

        self.assertEqual(_find_year_column(sheet, 1, 2026, "行业份额"), "C")


if __name__ == "__main__":
    unittest.main()
