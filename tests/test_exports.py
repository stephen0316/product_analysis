from io import BytesIO
import unittest

from openpyxl import load_workbook

from post_evaluation.constants import METRICS, PRODUCTS
from post_evaluation.exports import _EXPORT_TEMPLATE, make_result_tables, make_workbook
from post_evaluation.model_defaults import default_model_definition


class WorkbookExportTests(unittest.TestCase):
    def test_template_workbook_groups_indicators_by_business_dimension(self):
        scores = {}
        traces = []
        summary_rows = []
        for index, product in enumerate(PRODUCTS, 1):
            score = 10 - (index - 1) * 0.5
            reverse_rank = len(PRODUCTS) - index + 1
            scores[product] = {key: score for key, _, _ in METRICS}
            traces.append({
                "product": product, "metric": "market_share", "rank": reverse_rank, "score": score,
                "steps": [{"label": "2025年市场份额", "value": index / 100, "rank": reverse_rank, "score": score}],
            })
            for metric in ("revenue", "gross_profit", "gross_margin", "operating_profit", "operating_margin"):
                previous = index * 100
                current = previous * 1.1
                metric_rank = reverse_rank if metric == "revenue" else index
                traces.append({
                    "product": product, "metric": metric, "rank": metric_rank, "score": score, "average": score,
                    "steps": [
                        {"label": "2024年原值", "value": previous},
                        {"label": "2025年原值", "value": current, "rank": metric_rank, "score": score},
                        {"label": "增幅", "value": 0.1, "rank": metric_rank, "score": score},
                        {"label": "两项得分平均", "value": score, "rank": metric_rank, "score": score},
                    ],
                })
            traces.append({
                "product": product, "metric": "three_eff", "rank": index, "score": score, "average": score,
                "steps": [
                    {"label": "劳产", "value": index * 10, "rank": index, "score": score},
                    {"label": "元效", "value": index * 20, "rank": index, "score": score},
                    {"label": "人效", "value": index * 30, "rank": index, "score": score},
                    {"label": "三项得分平均", "value": score, "rank": index, "score": score},
                ],
            })
            self_steps = [
                {"label": "自研收入占比", "value": 0.5, "rank": index, "score": score},
                {"label": "研发收入", "value": 1000 + index},
                *[{"label": f"研发成本{number}", "value": 10} for number in range(1, 8)],
                {"label": "研发成本合计", "value": 70},
                {"label": "研发投入产出比", "value": 14.3, "rank": index, "score": score},
                {"label": "两项得分平均", "value": score, "rank": index, "score": score},
            ]
            traces.append({"product": product, "metric": "self_research", "rank": reverse_rank, "score": score, "steps": self_steps})
            traces.append({
                "product": product, "metric": "ability", "rank": index, "score": score,
                "steps": [{"label": "模拟能力测评值", "value": 100 - index, "rank": index, "score": score}],
            })
            traces.append({
                "product": product, "metric": "quality", "rank": index, "score": score,
                "steps": [{"label": "2025年品质服务", "value": 95 - index / 10, "rank": index, "score": score}],
            })
            definition = default_model_definition()
            summary_rows.append({
                "product": product,
                "total": score * 10,
                "rank": reverse_rank,
                "complete": True,
                "dimensions": [
                    {"key": item["key"], "label": item["label"], "score": score, "missing": []}
                    for item in definition["dimensions"]
                ],
            })

        run = {"year": 2025, "data_year": 2025, "model": {"definition": default_model_definition()}}
        result = {
            "scores": scores,
            "traces": traces,
            "model_summary": {"ready": True, "rows": summary_rows},
        }
        workbook = load_workbook(BytesIO(make_workbook(run, result)), data_only=False)
        tables = make_result_tables(run, result)
        template = load_workbook(_EXPORT_TEMPLATE, data_only=False, keep_links=False)

        self.assertEqual(workbook.sheetnames, template.sheetnames)
        for title in workbook.sheetnames:
            self.assertEqual(workbook[title].freeze_panes, "C3")
            self.assertEqual(workbook[title]["A1"].font.bold, template["重点产品-产品三效"]["A1"].font.bold)
            self.assertEqual(workbook[title]["A1"].alignment.horizontal, template["重点产品-产品三效"]["A1"].alignment.horizontal)

        market = workbook["重点产品-市场份额"]
        market_table = next(table for table in tables if table["label"] == "市场份额")
        self.assertEqual(
            [column["label"] for column in market_table["columns"]],
            [market.cell(2, column).value for column in range(1, market.max_column + 1)],
        )
        self.assertEqual(market_table["rows"][0]["values"], [market.cell(3, column).value for column in range(1, market.max_column + 1)])
        self.assertEqual(market["B3"].value, PRODUCTS[-1])
        self.assertEqual([market.cell(2, column).value for column in range(1, 6)], [
            "序号", "产品线", "2025年市场份额", "2025年市场份额排名", "2025年市场份额得分",
        ])
        self.assertEqual(market["C3"].value, 0.13)
        self.assertEqual(market["D3"].value, 1)

        benefit = workbook["重点产品-效益贡献"]
        self.assertEqual(benefit["B3"].value, PRODUCTS[-1])
        benefit_headers = [benefit.cell(2, column).value for column in range(1, benefit.max_column + 1)]
        self.assertIn("2024年收入", benefit_headers)
        for label in ("收入", "毛利", "毛利率", "营业利润", "营业利润率"):
            self.assertIn(f"{label}/增幅平均分", benefit_headers)
            self.assertIn(f"{label}/增幅平均分排名", benefit_headers)
            self.assertIn(f"{label}/增幅平均分得分", benefit_headers)
        self.assertFalse(any("｜" in str(header) for header in benefit_headers))

        research = workbook["重点产品-自研能力"]
        research_headers = [research.cell(2, column).value for column in range(1, research.max_column + 1)]
        self.assertIn("研发总成本", research_headers)
        self.assertIn("自研能力平均分", research_headers)
        self.assertIn("自研能力平均分排名", research_headers)
        self.assertIn("自研能力平均分得分", research_headers)
        self.assertFalse(any(str(header).startswith("研发成本") and header != "研发总成本" for header in research_headers))
        self.assertEqual(research_headers.count("研发总成本"), 1)

        total = workbook["重点产品-总分"]
        total_table = tables[0]
        self.assertEqual(total_table["kind"], "total")
        self.assertEqual(
            [column["label"] for column in total_table["columns"]],
            [total.cell(2, column).value for column in range(1, total.max_column + 1)],
        )
        self.assertEqual(total_table["rows"][0]["values"], [total.cell(3, column).value for column in range(1, total.max_column + 1)])
        self.assertEqual(total["B3"].value, PRODUCTS[-1])
        self.assertEqual(total["M3"].value, 40)
        self.assertEqual(total["N3"].value, 1)

        custom_definition = default_model_definition()
        custom_definition["categories"] = [
            {"key": "profit", "label": "盈利能力", "export_sort_indicator_key": "economic_gross_profit"},
            {"key": "competition", "label": "市场竞争"},
        ]
        selected = []
        for indicator in custom_definition["dimensions"]:
            if indicator["key"] in {"economic_revenue", "economic_gross_profit"}:
                selected.append({**indicator, "category_key": "profit"})
            elif indicator["key"] == "market":
                selected.append({**indicator, "label": "行业份额", "category_key": "competition"})
        custom_definition["dimensions"] = selected
        custom = load_workbook(BytesIO(make_workbook({"year": 2025, "data_year": 2025, "model": {"definition": custom_definition}}, result)))
        custom_tables = make_result_tables({"year": 2025, "data_year": 2025, "model": {"definition": custom_definition}}, result)
        self.assertEqual(custom.sheetnames, ["重点产品-雷达图维度", "重点产品-盈利能力", "重点产品-市场竞争", "重点产品-总分"])
        self.assertEqual([table["label"] for table in custom_tables], ["总分", "盈利能力", "市场竞争"])
        self.assertEqual(custom["重点产品-盈利能力"]["B3"].value, PRODUCTS[0])
        self.assertEqual(
            [custom["重点产品-市场竞争"].cell(2, column).value for column in range(1, 6)],
            ["序号", "产品线", "2025年行业份额", "2025年行业份额排名", "2025年行业份额得分"],
        )
        self.assertEqual(custom["重点产品-总分"]["C2"].value, "行业份额")
        self.assertNotIn("重点产品-效益贡献", custom.sheetnames)

        future = load_workbook(BytesIO(make_workbook({"year": 2026, "data_year": 2026, "model": {"definition": custom_definition}}, result)))
        self.assertEqual(future["重点产品-市场竞争"]["C2"].value, "2026年行业份额")
        self.assertIn("2026年", future["重点产品-总分"]["A1"].value)


if __name__ == "__main__":
    unittest.main()
