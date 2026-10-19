from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from post_evaluation.app import app
from post_evaluation.constants import PRODUCTS


SOURCE_DIR = Path("/Users/apple/Desktop/后评估")
SOURCES = (
    (["market_share"], "产品市场份额信息统计0113.xlsx"),
    (["revenue", "gross_profit", "gross_margin", "operating_profit", "operating_margin"], "产品后评估-财务部-原表.xlsx"),
    (["three_eff"], "20260728-2526产品线三效数据-发市场部（定稿）-补充2506三效.xlsx"),
)
SELF_RESEARCH_SOURCES = (
    "25年底产品线自有占比0508-【经分】.xlsx",
    "20260721-202512研发外采数据-财务部.xlsx",
)


def dimension_targets(run, metrics):
    return [
        dimension["key"]
        for dimension in run["model"]["definition"]["dimensions"]
        if set(metrics).intersection(dimension.get("metric_keys", []))
    ]


@unittest.skipUnless(all((SOURCE_DIR / name).exists() for _, name in SOURCES), "本机样本附件不可用")
class EndToEndTests(unittest.TestCase):
    def test_two_self_research_files_match_the_same_dimension(self):
        if not all((SOURCE_DIR / name).exists() for name in SELF_RESEARCH_SOURCES):
            self.skipTest("本机自研能力样本附件不可用")
        with tempfile.TemporaryDirectory() as directory:
            old = os.environ.get("PRODUCT_EVAL_DATA_DIR")
            os.environ["PRODUCT_EVAL_DATA_DIR"] = directory
            try:
                client = TestClient(app)
                run = client.post("/api/runs", json={"year": 2025, "name": "自研能力多来源验证"}).json()
                self.assertEqual(run["data_year"], 2025)
                expected_targets = dimension_targets(run, ["self_research"])
                opened = [(name, (SOURCE_DIR / name).open("rb")) for name in SELF_RESEARCH_SOURCES]
                try:
                    response = client.post(
                        f"/api/runs/{run['id']}/files",
                        files=[("files", (name, source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for name, source in opened],
                    )
                finally:
                    for _, source in opened:
                        source.close()
                self.assertEqual(response.status_code, 200, response.text)
                uploaded = response.json()["files"]
                self.assertEqual([file["targets"] for file in uploaded], [expected_targets, expected_targets])
                self.assertEqual(uploaded[0]["mapping"]["sheet"], "Sheet1")
                self.assertEqual(uploaded[1]["mapping"]["sheet"], "研发V2")
                self.assertEqual(uploaded[0]["mapping"]["provided_summary"], ["2025年自研收入占比"])
                self.assertEqual(
                    uploaded[1]["mapping"]["provided_summary"],
                    ["2025年研发收入", "2025年研发成本（E～K项合计）"],
                )
                self.assertEqual(uploaded[1]["mapping"]["fields"]["research_revenue:2025"], "C")
                self.assertEqual(
                    [uploaded[1]["mapping"]["fields"][f"{key}:2025"] for key in (
                        "research_cost_own_labor", "research_cost_partner_labor", "research_cost_travel",
                        "research_cost_external_independent", "research_cost_external_internal",
                        "research_cost_resource_pool", "research_cost_allocation",
                    )],
                    ["E", "F", "G", "H", "I", "J", "K"],
                )
                validated = client.post(f"/api/runs/{run['id']}/validate")
                self.assertEqual(validated.status_code, 200, validated.text)
                self.assertEqual(validated.json()["validation"]["status"], "passed")
                calculated = client.post(f"/api/runs/{run['id']}/calculate")
                self.assertEqual(calculated.status_code, 200, calculated.text)
                result = calculated.json()
                self.assertEqual(next(item for item in result["metrics"] if item["key"] == "self_research")["status"], "已计算")
                self.assertEqual(len([trace for trace in result["traces"] if trace["metric"] == "self_research"]), len(PRODUCTS))
                self.assertFalse([issue for issue in result["issues"] if issue["level"] == "error"])
            finally:
                if old is None:
                    os.environ.pop("PRODUCT_EVAL_DATA_DIR", None)
                else:
                    os.environ["PRODUCT_EVAL_DATA_DIR"] = old

    def test_multiple_files_can_match_the_same_dimension(self):
        with tempfile.TemporaryDirectory() as directory:
            old = os.environ.get("PRODUCT_EVAL_DATA_DIR")
            os.environ["PRODUCT_EVAL_DATA_DIR"] = directory
            try:
                client = TestClient(app)
                run_id = client.post("/api/runs", json={"year": 2025, "name": "来源唯一性验证"}).json()["id"]
                expected_targets = dimension_targets(client.get(f"/api/runs/{run_id}").json(), ["market_share"])
                source_path = SOURCE_DIR / "产品市场份额信息统计0113.xlsx"
                for _ in range(2):
                    with source_path.open("rb") as source:
                        response = client.post(f"/api/runs/{run_id}/files", files={"files": (source_path.name, source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
                    self.assertEqual(response.status_code, 200, response.text)
                    uploaded = response.json()["files"][-1]
                    self.assertEqual(uploaded["targets"], expected_targets)
                self.assertEqual(uploaded["mapping"]["suggested_targets"], expected_targets)
            finally:
                if old is None:
                    os.environ.pop("PRODUCT_EVAL_DATA_DIR", None)
                else:
                    os.environ["PRODUCT_EVAL_DATA_DIR"] = old

    def test_delete_uploaded_file_and_upload_again(self):
        with tempfile.TemporaryDirectory() as directory:
            old = os.environ.get("PRODUCT_EVAL_DATA_DIR")
            os.environ["PRODUCT_EVAL_DATA_DIR"] = directory
            try:
                client = TestClient(app)
                run_id = client.post("/api/runs", json={"year": 2025, "name": "文件删除验证"}).json()["id"]
                expected_targets = dimension_targets(client.get(f"/api/runs/{run_id}").json(), ["market_share"])
                source_path = SOURCE_DIR / "产品市场份额信息统计0113.xlsx"
                with source_path.open("rb") as source:
                    uploaded = client.post(f"/api/runs/{run_id}/files", files={"files": (source_path.name, source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
                self.assertEqual(uploaded.status_code, 200, uploaded.text)
                file_id = uploaded.json()["files"][0]["id"]
                deleted = client.delete(f"/api/runs/{run_id}/files/{file_id}")
                self.assertEqual(deleted.status_code, 200, deleted.text)
                self.assertEqual(deleted.json()["files"], [])
                with source_path.open("rb") as source:
                    uploaded_again = client.post(f"/api/runs/{run_id}/files", files={"files": (source_path.name, source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
                self.assertEqual(uploaded_again.status_code, 200, uploaded_again.text)
                self.assertEqual(uploaded_again.json()["files"][0]["targets"], expected_targets)
            finally:
                if old is None:
                    os.environ.pop("PRODUCT_EVAL_DATA_DIR", None)
                else:
                    os.environ["PRODUCT_EVAL_DATA_DIR"] = old

    def test_upload_calculate_and_export(self):
        with tempfile.TemporaryDirectory() as directory:
            old = os.environ.get("PRODUCT_EVAL_DATA_DIR")
            os.environ["PRODUCT_EVAL_DATA_DIR"] = directory
            try:
                client = TestClient(app)
                created = client.post("/api/runs", json={"year": 2025, "name": "自动化验证"})
                self.assertEqual(created.status_code, 200)
                self.assertEqual(created.json()["data_year"], 2025)
                run_id = created.json()["id"]
                current_run = client.get(f"/api/runs/{run_id}").json()
                opened = [(targets, name, (SOURCE_DIR / name).open("rb")) for targets, name in SOURCES]
                try:
                    response = client.post(
                        f"/api/runs/{run_id}/files",
                        files=[("files", (name, source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for _, name, source in opened],
                    )
                finally:
                    for _, _, source in opened:
                        source.close()
                self.assertEqual(response.status_code, 200, response.text)
                uploaded = response.json()["files"]
                self.assertEqual(len(uploaded), len(SOURCES))
                for file, (metrics, _) in zip(uploaded, SOURCES):
                    self.assertIsNotNone(file["mapping"]["sheet"])
                    self.assertEqual(file["targets"], dimension_targets(current_run, metrics))
                market_mapping = uploaded[0]["mapping"]
                self.assertEqual(market_mapping["fields"], {"market_share:2025": "G"})
                validated = client.post(f"/api/runs/{run_id}/validate")
                self.assertEqual(validated.status_code, 200, validated.text)
                self.assertEqual(validated.json()["validation"]["status"], "passed")
                response = client.post(f"/api/runs/{run_id}/calculate")
                self.assertEqual(response.status_code, 200, response.text)
                result = response.json()
                self.assertEqual(result["covered_points"], 80)
                self.assertFalse(result["official_total_available"])
                self.assertEqual(len(result["traces"]), len(PRODUCTS) * 8)
                self.assertFalse([i for i in result["issues"] if i["level"] == "error"])
                self.assertEqual([m["status"] for m in result["metrics"]][-3:], ["待完成", "已计算", "待完成"])
                self.assertTrue(all(len(result["scores"][p]) == 8 for p in PRODUCTS))

                workbook_response = client.get(f"/api/runs/{run_id}/export.xlsx")
                self.assertEqual(workbook_response.status_code, 200)
                wb = load_workbook(BytesIO(workbook_response.content))
                self.assertEqual(wb.sheetnames, [
                    "重点产品-雷达图维度",
                    "重点产品-市场份额",
                    "重点产品-效益贡献",
                    "重点产品-产品三效",
                    "重点产品-自研能力",
                    "重点产品-品质服务",
                    "重点产品-总分",
                ])
                ws = wb["重点产品-总分"]
                self.assertEqual(ws["C3"].value, result["scores"]["OnePark"]["market_share"])
                self.assertEqual(ws["L3"].value, "待完成")
                self.assertEqual(ws["M3"].value, "待完成")
                self.assertEqual(ws["N3"].value, "待完成")
                market = wb["重点产品-市场份额"]
                self.assertEqual([market.cell(2, column).value for column in range(3, 6)], [
                    "2025年市场份额", "2025年市场份额排名", "2025年市场份额得分",
                ])
                market_rows = {market.cell(row, 2).value: row for row in range(3, market.max_row + 1)}
                card_row = market_rows["物联卡"]
                self.assertIsInstance(market.cell(card_row, 3).value, (int, float))
                self.assertIsInstance(market.cell(card_row, 4).value, (int, float))
                self.assertEqual(wb["重点产品-产品三效"]["C2"].value, "2025年劳产")
                research_headers = [wb["重点产品-自研能力"].cell(2, column).value for column in range(1, wb["重点产品-自研能力"].max_column + 1)]
                self.assertIn("自研能力得分", research_headers)
                self.assertIn("2025年能力评测（模拟值）得分", research_headers)
                self.assertEqual(wb["重点产品-品质服务"]["C3"].value, "待完成")

                self.assertEqual(client.get(f"/api/runs/{run_id}/report.docx").status_code, 409)
                generated = client.post(f"/api/runs/{run_id}/report", json={})
                self.assertEqual(generated.status_code, 200, generated.text)
                self.assertIsNotNone(generated.json()["report"])
                report_response = client.get(f"/api/runs/{run_id}/report.docx")
                self.assertEqual(report_response.status_code, 200)
                document = Document(BytesIO(report_response.content))
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn("不生成正式总分", text)
                self.assertEqual(len(document.inline_shapes), 1 + len(PRODUCTS))
            finally:
                if old is None:
                    os.environ.pop("PRODUCT_EVAL_DATA_DIR", None)
                else:
                    os.environ["PRODUCT_EVAL_DATA_DIR"] = old


if __name__ == "__main__":
    unittest.main()
