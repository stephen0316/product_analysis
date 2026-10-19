from base64 import b64decode
from copy import deepcopy
from io import BytesIO
import json
import os
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from post_evaluation import store
from post_evaluation.app import app
from post_evaluation.constants import PRODUCTS
from post_evaluation.exports import make_report
from post_evaluation.gemini_rules import GeminiNotConfigured, GeminiRuleError
from post_evaluation.model_defaults import default_model_definition
from post_evaluation.report_analysis import _generate_product, generate_evaluation_draft, report_evidence, validate_draft
from test_execution_boundary import calculate_sample


class ReportAnalysisTests(unittest.TestCase):
    def setUp(self):
        gateway = patch.dict(os.environ, LLM_API_KEY="test-secret", GEMINI_API_KEY="",
                             LLM_MODEL="test-model", LLM_REPORT_MODEL="report-model")
        gateway.start()
        self.addCleanup(gateway.stop)

    @classmethod
    def setUpClass(cls):
        cls.definition = default_model_definition()
        cls.definition["dimensions"][0]["label"] = "行业份额"
        cls.report_run = {"year": 2025, "data_year": 2025, "model": {"definition": cls.definition},
                   "notes": {}, "files": []}
        cls.result = calculate_sample(cls.definition)

    def draft(self, evidence):
        return {"products": [{"product": row["product"], "evaluation": "已提供数据反映了当前指标之间的差异。",
                              "recommendation": "建议复核相对较弱指标的经营原因并跟踪改善效果。",
                              "evidence_ids": [row["facts"][0]["id"]]} for row in evidence["products"]]}

    def test_evidence_uses_current_headers_and_excludes_simulated_metrics(self):
        evidence = report_evidence(self.report_run, self.result)
        self.assertEqual(len(evidence["products"]), 13)
        self.assertTrue(any("行业份额" in fact["label"] for fact in evidence["products"][0]["facts"]))
        for row in evidence["products"]:
            self.assertFalse(any(fact["metric"] == "ability" for fact in row["facts"]))
            self.assertFalse(any("source_name" in fact or "file_id" in fact for fact in row["facts"]))
        changed = deepcopy(self.result)
        changed["traces"][0]["steps"][0]["value"] += .1
        self.assertNotEqual(evidence["signature"], report_evidence(self.report_run, changed)["signature"])

    def test_rejects_incomplete_duplicate_and_cross_product_citations(self):
        evidence = report_evidence(self.report_run, self.result)
        draft = self.draft(evidence)
        validate_draft(draft, evidence)
        broken = deepcopy(draft)
        broken["products"][0]["evidence_ids"] = broken["products"][1]["evidence_ids"]
        with self.assertRaises(ValueError): validate_draft(broken, evidence)

    def test_prose_cleanup_preserves_numbers_metadata_and_original_draft(self):
        evidence = report_evidence(self.report_run, self.result)
        draft = self.draft(evidence)
        draft["products"][0]["evaluation"] = "平均分为7.0，排名第七P1D3C11|P1D3C12。自研得分10.0【P1D4C7】。"
        draft["products"][0]["recommendation"] = "建议核实收入下滑原因（P1D2C3|P1D2C4），跟踪2025年收入。"
        original = deepcopy(draft)
        cleaned = validate_draft(draft, evidence)
        self.assertEqual(cleaned["products"][0]["evaluation"], "平均分为7.0，排名第七。自研得分10.0。")
        self.assertEqual(cleaned["products"][0]["recommendation"], "建议核实收入下滑原因，跟踪2025年收入。")
        self.assertEqual(cleaned["products"][0]["evidence_ids"], original["products"][0]["evidence_ids"])
        self.assertEqual(draft, original)
        broken = deepcopy(draft)
        broken["products"][0]["evaluation"] = "P1D3C11|P1D3C12"
        with self.assertRaisesRegex(ValueError, "缺少正文"):
            validate_draft(broken, evidence)
        with self.assertRaises(ValueError): validate_draft({"products": draft["products"][:-1]}, evidence)
        broken = deepcopy(draft)
        broken["products"][0]["product"] = broken["products"][1]["product"]
        with self.assertRaises(ValueError): validate_draft(broken, evidence)

    def test_ai_call_validates_output_and_redacts_errors(self):
        evidence = report_evidence(self.report_run, self.result)
        payloads = []
        def provider(request, **kwargs):
            payload = json.loads(request.data)
            payloads.append(payload)
            scope = json.loads(payload["messages"][1]["content"])
            row = self.draft(scope)["products"][0]
            return BytesIO(json.dumps({"choices": [{"finish_reason": "STOP", "message": {"content": json.dumps(row)}}]}).encode())
        with patch.dict(os.environ, LLM_API_KEY="test-secret"), patch("post_evaluation.llm_client.urlopen", side_effect=provider):
            draft = generate_evaluation_draft(self.report_run, self.result)
            self.assertEqual(draft["signature"], evidence["signature"])
            self.assertEqual(len(payloads), len(PRODUCTS))
            for payload in payloads:
                self.assertEqual(payload["model"], "report-model")
                text = payload["messages"][1]["content"]
                self.assertNotIn("source_name", text)
                scope = json.loads(text)
                self.assertEqual(len(scope["products"]), 1)
                schema = json.loads(payload["messages"][0]["content"].rsplit("\n", 1)[1])
                self.assertEqual(schema["properties"]["product"]["enum"], [scope["products"][0]["product"]])
                self.assertEqual(set(schema["properties"]["evidence_ids"]["items"]["enum"]), {f["id"] for f in scope["products"][0]["facts"]})
        with patch.dict(os.environ, LLM_API_KEY=""):
            with self.assertRaises(GeminiNotConfigured): generate_evaluation_draft(self.report_run, self.result)

    def test_retry_repairs_bad_reference_without_weakening_validation(self):
        evidence = report_evidence(self.report_run, self.result)
        product = evidence["products"][0]
        row = self.draft({"products": [product]})["products"][0]
        bad = {**row, "evidence_ids": ["P1P1D2C31"]}
        def response(value):
            return BytesIO(json.dumps({"choices": [{"finish_reason": "STOP", "message": {"content": json.dumps(value)}}]}).encode())
        with patch("post_evaluation.llm_client.urlopen", side_effect=[response(bad), response(row)]) as request:
            result = _generate_product(product, evidence, "test-model", "test-secret")
            self.assertEqual(result["evidence_ids"], row["evidence_ids"])
            self.assertEqual(request.call_count, 2)
            self.assertEqual(json.loads(request.call_args.args[0].data)["model"], "test-model")
        with patch("post_evaluation.llm_client.urlopen", side_effect=[response(bad), response(bad)]):
            with self.assertRaisesRegex(GeminiRuleError, "数据依据"):
                _generate_product(product, evidence, "test-model", "test-secret")

    def test_truncation_is_distinguished_from_transport_errors(self):
        evidence = report_evidence(self.report_run, self.result)
        product = evidence["products"][0]
        output = json.dumps({"choices": [{"finish_reason": "length"}]}).encode()
        with patch("post_evaluation.llm_client.urlopen", side_effect=[BytesIO(output), BytesIO(output)]):
            with self.assertRaisesRegex(GeminiRuleError, "截断"):
                _generate_product(product, evidence, "test-model", "test-secret")

    def test_report_contains_all_products_dynamic_dimensions_and_edited_draft(self):
        evidence = report_evidence(self.report_run, self.result)
        draft = {**self.draft(evidence), "signature": evidence["signature"]}
        draft["products"][0]["evaluation"] = "指标得分8.0P1D1C2|P1D1C3，需继续跟踪【P1D2C4】。"
        draft["products"][0]["recommendation"] = "已人工编辑的具体提升建议（P1D1C2）。"
        original = deepcopy(draft)
        chart_bytes = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jr1sAAAAASUVORK5CYII=")
        with patch("post_evaluation.report_sections.product_radar", side_effect=lambda _: BytesIO(chart_bytes)):
            document = Document(BytesIO(make_report({**self.report_run, "notes": {"evaluation_draft": draft}}, self.result)))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("附录 重点产品详细评价和提升建议", text)
        self.assertIn("行业份额", text)
        self.assertIn("已人工编辑的具体提升建议", text)
        self.assertIn("指标得分8.0，需继续跟踪。", text)
        self.assertNotRegex(text, r"P\d+D\d+C\d+")
        self.assertFalse(any(p.text.startswith(("业务数据：", "评价依据：")) for p in document.paragraphs))
        self.assertEqual(draft, original)
        paragraphs = document.paragraphs
        for index, paragraph in enumerate(paragraphs):
            if paragraph.text.startswith("评分表现："):
                self.assertTrue(paragraphs[index + 1]._p.xpath(".//w:drawing"))
                self.assertTrue(paragraphs[index + 4].text.startswith("详细评价："))
        headings = [p.text for p in document.paragraphs if p.style.name == "Heading 3"]
        self.assertEqual(len(headings), len(PRODUCTS))
        self.assertTrue(all(any(p in heading for heading in headings) for p in PRODUCTS))

    def test_generation_does_not_persist_draft_and_rejects_stale_evidence(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, PRODUCT_EVAL_DATA_DIR=directory):
            run = store.create_run(2025, "报告流程")
            store.update_run_json(run["id"], "result", self.result)
            client = TestClient(app)
            path = f"/api/runs/{run['id']}"
            evidence = report_evidence(store.get_run(run["id"]), self.result)
            draft = {**self.draft(evidence), "signature": evidence["signature"]}
            draft["products"][0]["evaluation"] = "已提供数据反映了当前指标之间的差异P1D1C2。"
            with patch("post_evaluation.app.generate_evaluation_draft", return_value=draft):
                self.assertEqual(client.post(path + "/report-analysis", json={}).status_code, 200)
            self.assertIsNone(store.get_run(run["id"])["report"])
            self.assertNotIn("evaluation_draft", store.get_run(run["id"])["notes"])
            with patch("post_evaluation.app.make_report", return_value=b"snapshot"):
                response = client.post(path + "/report", json={"evaluation_draft": draft})
                self.assertEqual(response.status_code, 200, response.text)
            saved_draft = store.get_run(run["id"])["notes"]["evaluation_draft"]
            self.assertEqual(saved_draft["products"][0]["evaluation"], "已提供数据反映了当前指标之间的差异。")
            self.assertEqual(saved_draft["products"][0]["evidence_ids"], draft["products"][0]["evidence_ids"])
            response = client.post(path + "/report", json={"analysis": "新说明", "evaluation_draft": draft})
            self.assertEqual(response.status_code, 409)
            self.assertEqual(client.get(path + "/report.docx").content, b"snapshot")

    def test_optional_supplements_omit_blanks_and_preserve_legacy_review(self):
        with patch("post_evaluation.report_sections.product_radar", return_value=None):
            empty = Document(BytesIO(make_report({**self.report_run, "notes": {
                "analysis": "  ", "recommendations": "", "review": "\n",
            }}, self.result)))
            filled = Document(BytesIO(make_report({**self.report_run, "notes": {
                "analysis": "已核实的经营背景", "recommendations": " ", "review": "原审阅意见内容",
            }}, self.result)))
        empty_text = "\n".join(p.text for p in empty.paragraphs)
        for label in ("三 补充说明", "经营原因分析", "产品处置建议", "报告备注", "待负责人填写并确认。"):
            self.assertNotIn(label, empty_text)
        filled_text = "\n".join(p.text for p in filled.paragraphs)
        for label in ("三 补充说明", "经营原因分析", "已核实的经营背景", "报告备注", "原审阅意见内容"):
            self.assertIn(label, filled_text)
        self.assertNotIn("产品处置建议", filled_text)
        self.assertNotIn("负责人审阅意见", filled_text)
        evidence = report_evidence({**self.report_run, "notes": {
            "analysis": "经营背景", "recommendations": "已有处置方向", "review": "仅报告备注",
        }}, self.result)
        self.assertEqual(evidence["business_context"], {"analysis": "经营背景", "recommendations": "已有处置方向"})
