from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .constants import (
    ECONOMIC_FIELDS,
    PRODUCTS,
    SELF_RESEARCH_COST_FIELDS,
    adapter_for_metrics,
    mapping_fields_for_metrics,
)


def normalize_product(value: Any, role: str) -> str | None:
    if value is None:
        return None
    compact = re.sub(r"\s+", "", str(value)).casefold()
    aliases = {re.sub(r"\s+", "", p).casefold(): p for p in PRODUCTS}
    aliases.update({
        "onenet城市物联网平台": "OneNET",
        "onepark智慧园区": "OnePark",
        "onesec物联网安全解决方案": "OneSec",
    })
    if role == "market":
        aliases.update({"安全芯片": "芯片", "芯片（安全)": "芯片", "芯片(安全)": "芯片"})
    if role == "self_research":
        # Settlement rows do not participate, except for the explicitly
        # confirmed IoT-card exception.
        if compact == "物联网卡结算":
            return "物联卡"
        if "结算" in compact:
            return None
    return aliases.get(compact)


def inspect_workbook(path: str | Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    result = []
    for ws in wb:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 7), max_col=min(ws.max_column, 25), values_only=True):
            preview.append([display_value(v) for v in row])
        result.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "preview": preview})
    wb.close()
    return result


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.8g}"
    return str(value)[:160]


def auto_mapping(path: str | Path, metric_keys: list[str], year: int) -> dict[str, Any]:
    """Suggest a mapping for the dimensions the user assigned to this file.

    ``role`` is retained in stored records as a reader hint only; it is no
    longer selected by the business user.
    """
    role = adapter_for_metrics(metric_keys)
    wanted_fields = set(mapping_fields_for_metrics(metric_keys, year))
    wb = load_workbook(path, read_only=True, data_only=False)
    candidate = None
    for ws in wb:
        first = " ".join(str(ws.cell(r, c).value or "") for r in range(1, min(4, ws.max_row) + 1) for c in range(1, min(ws.max_column, 14) + 1))
        if role == "market" and any(label in first for label in ("市场份额", "行业份额")) and "公司产品" in first:
            candidate = ws
            share_label = "行业份额" if "行业份额" in first else "市场份额"
            header_row = next((
                row for row in range(1, min(ws.max_row, 5) + 1)
                if any("公司产品" in str(ws.cell(row, column).value or "") for column in range(1, ws.max_column + 1))
                and any(share_label in str(ws.cell(row, column).value or "") for column in range(1, ws.max_column + 1))
            ), 1)
            product_col = next((
                get_column_letter(column)
                for column in range(1, ws.max_column + 1)
                if "公司产品" in str(ws.cell(header_row, column).value or "")
            ), "D")
            mapping = {
                "sheet": ws.title,
                "header_row": header_row,
                "product_col": product_col,
                "fields": {f"market_share:{year}": _find_year_column(ws, header_row, year, share_label)},
            }
            break
        if role == "finance" and "产品收入" in first and "产品毛利" in first:
            finance_mapping = _find_finance_mapping(ws, year)
            if not finance_mapping:
                continue
            candidate = ws
            mapping = finance_mapping
            break
        if role == "three_eff" and all(word in first for word in ("劳产", "元效", "人效")):
            candidate = ws
            header_row = next((r for r in range(1, min(ws.max_row, 5) + 1) if any("劳产" in str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))), 2)
            def header_col(label: str) -> str | None:
                return next((get_column_letter(c) for c in range(1, ws.max_column + 1) if label in str(ws.cell(header_row, c).value or "")), None)
            mapping = {"sheet": ws.title, "header_row": header_row, "product_col": header_col("产品线") or "B", "fields": {f"labor_output:{year}": header_col("劳产"), f"yuan_efficiency:{year}": header_col("元效"), f"human_efficiency:{year}": header_col("人效")}}
            break
        if "self_research" in metric_keys and "自有占比" in first and "产品线" in first:
            candidate = ws
            header_row = next((r for r in range(1, min(ws.max_row, 5) + 1) if any("自有占比" in str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))), 1)
            def self_owned_col(label: str) -> str | None:
                return next((get_column_letter(c) for c in range(1, ws.max_column + 1) if label in str(ws.cell(header_row, c).value or "")), None)
            mapping = {
                "sheet": ws.title,
                "header_row": header_row,
                "product_col": self_owned_col("产品线") or "B",
                "fields": {f"self_owned_ratio:{year}": self_owned_col("自有占比")},
            }
            break
        if "self_research" in metric_keys and "研发收入" in first and "研发成本" in first and "产品线" in first:
            candidate = ws
            header_row = next((r for r in range(1, min(ws.max_row, 5) + 1) if any("研发收入" in str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))), 2)
            product_col = next((get_column_letter(c) for c in range(1, ws.max_column + 1) if "产品线" in str(ws.cell(header_row, c).value or "")), "B")
            header_lookup = {
                re.sub(r"\s+", "", str(ws.cell(header_row, c).value or "")): get_column_letter(c)
                for c in range(1, ws.max_column + 1)
            }
            cost_labels = {
                "research_cost_own_labor": "研究开发费-自有人工-项目成本",
                "research_cost_partner_labor": "研究开发费-合作人工",
                "research_cost_travel": "研究开发费-差旅费",
                "research_cost_external_independent": "研究开发费-其他外采（自立）",
                "research_cost_external_internal": "研究开发费-其他外采（委内）",
                "research_cost_resource_pool": "研究开发费-资源池成本",
                "research_cost_allocation": "研究开发费-其他分摊-项目成本",
            }
            fields = {f"research_revenue:{year}": next((column for label, column in header_lookup.items() if "研发收入" in label), None)}
            for key in SELF_RESEARCH_COST_FIELDS:
                wanted = re.sub(r"\s+", "", cost_labels[key])
                fields[f"{key}:{year}"] = next((column for label, column in header_lookup.items() if label == wanted), None)
            mapping = {"sheet": ws.title, "header_row": header_row, "product_col": product_col, "fields": fields}
            break
    wb.close()
    if candidate is None:
        return {"sheet": None, "header_row": 1, "product_col": "A", "fields": {key: None for key in wanted_fields}}
    mapping["fields"] = {key: value for key, value in mapping["fields"].items() if key in wanted_fields}
    for key in wanted_fields:
        mapping["fields"].setdefault(key, None)
    return mapping


def detect_metrics(path: str | Path) -> list[str]:
    """Identify familiar source layouts without asking the user for a file type."""
    wb = load_workbook(path, read_only=True, data_only=False)
    detected: list[str] = []
    try:
        for ws in wb:
            first = " ".join(
                str(ws.cell(r, c).value or "")
                for r in range(1, min(4, ws.max_row) + 1)
                for c in range(1, min(ws.max_column, 14) + 1)
            )
            if any(label in first for label in ("市场份额", "行业份额")) and "公司产品" in first:
                detected.append("market_share")
            if "产品收入" in first and "产品毛利" in first:
                detected.extend(ECONOMIC_FIELDS)
            if all(word in first for word in ("劳产", "元效", "人效")):
                detected.append("three_eff")
            if "自有占比" in first and "产品线" in first:
                detected.append("self_research")
            if "研发收入" in first and "研发成本" in first and "产品线" in first:
                detected.append("self_research")
    finally:
        wb.close()
    return list(dict.fromkeys(detected))


def _find_year_column(ws, row: int, year: int, word: str) -> str | None:
    for c in range(1, ws.max_column + 1):
        label = str(ws.cell(row, c).value or "")
        if word in label and str(year) in label:
            return get_column_letter(c)
    return None


def _year_in_label(label: str, year: int) -> bool:
    compact = re.sub(r"\s+", "", label)
    return str(year) in compact or f"{year % 100:02d}年" in compact


def _find_finance_mapping(ws, year: int) -> dict[str, Any] | None:
    """Locate finance inputs by their two-row business headers and year labels."""
    group_row = next((
        row for row in range(1, min(ws.max_row, 6) + 1)
        if any("产品收入" in str(ws.cell(row, column).value or "") for column in range(1, ws.max_column + 1))
        and any("产品毛利" in str(ws.cell(row, column).value or "") for column in range(1, ws.max_column + 1))
    ), None)
    if group_row is None or group_row >= ws.max_row:
        return None
    detail_row = group_row + 1
    product_col = next((
        get_column_letter(column)
        for column in range(1, ws.max_column + 1)
        if "产品线" in str(ws.cell(group_row, column).value or "")
    ), "A")
    fields: dict[str, str | None] = {}
    active_group = ""
    for column in range(1, ws.max_column + 1):
        group_value = str(ws.cell(group_row, column).value or "").strip()
        if group_value:
            active_group = re.sub(r"\s+", "", group_value)
        detail = re.sub(r"\s+", "", str(ws.cell(detail_row, column).value or ""))
        for period in (year - 1, year):
            if not _year_in_label(detail, period):
                continue
            key = None
            if active_group == "产品收入" and "考核目标" not in detail:
                key = "revenue"
            elif active_group == "产品毛利":
                key = "gross_profit"
            elif active_group == "产品毛利率":
                key = "gross_margin"
            elif "营业利润" in active_group and "营业利润额" in detail:
                key = "operating_profit"
            elif "营业利润" in active_group and "营业利润率" in detail:
                key = "operating_margin"
            if key:
                fields[f"{key}:{period}"] = get_column_letter(column)
    return {
        "sheet": ws.title,
        "header_row": detail_row,
        "product_col": product_col,
        "fields": fields,
    }


def extract_records(path: str | Path, file_id: str, role: str, mapping: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    issues = []
    if not mapping.get("sheet"):
        return [], [{"level": "error", "code": "mapping", "message": "尚未确认工作表与字段映射"}]
    wb = load_workbook(path, read_only=True, data_only=True)
    if mapping["sheet"] not in wb:
        wb.close()
        return [], [{"level": "error", "code": "sheet", "message": "映射的工作表不存在"}]
    ws = wb[mapping["sheet"]]
    fields = mapping.get("fields", {})
    normalized_role = "self_research" if any(
        key.startswith(("self_owned_ratio:", "research_revenue:", "research_cost_"))
        for key in fields
    ) else role
    records = []
    seen = defaultdict(list)
    try:
        product_col = _column_number(mapping.get("product_col"))
        resolved = {key: _column_number(col) for key, col in fields.items() if col}
    except ValueError as exc:
        wb.close()
        return [], [{"level": "error", "code": "column", "message": str(exc)}]
    for row in range(int(mapping.get("header_row", 1)) + 1, ws.max_row + 1):
        original = ws.cell(row, product_col).value
        product = normalize_product(original, normalized_role)
        if not product:
            continue
        for key, col in resolved.items():
            cell = ws.cell(row, col)
            value = cell.value
            if value is None or value == "":
                continue
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                issues.append({"level": "error", "code": "nonnumeric", "message": f"{product} 的 {key} 不是数值", "field": key, "product": product, "cell": f"{ws.title}!{cell.coordinate}"})
                continue
            record = {"product": product, "field": key, "value": float(value), "file_id": file_id, "cell": f"{ws.title}!{cell.coordinate}", "source_name": str(original), "unit": mapping.get("units", {}).get(key, "")}
            records.append(record)
            seen[(product, key)].append(record)
    wb.close()
    for (product, key), matches in seen.items():
        if len(matches) > 1:
            issues.append({"level": "error", "code": "duplicate", "message": f"同一文件中 {product} 的 {key} 出现 {len(matches)} 次", "field": key, "product": product, "cells": [r["cell"] for r in matches]})
    return records, issues


def _column_number(value: Any) -> int:
    from openpyxl.utils import column_index_from_string

    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, str) and re.fullmatch(r"[A-Za-z]{1,3}", value.strip()):
        return column_index_from_string(value.strip().upper())
    raise ValueError(f"无效列号：{value}")
