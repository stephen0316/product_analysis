"""Narrative sections distilled from the supplied historical report."""
from io import BytesIO
import math
from pathlib import Path

from docx.shared import Inches
import matplotlib.pyplot as plt
from matplotlib import font_manager

from .constants import PRODUCTS
from .report_text import narrative_text


def real_scores(result: dict, product: str) -> list[dict]:
    simulated = set(result.get("simulated_metrics") or [])
    return [row for row in result.get("indicator_scores", {}).get(product, [])
            if row.get("ratio") is not None and not simulated.intersection(row.get("metric_keys", []))]


def add_dimension_analysis(doc, run: dict, result: dict) -> None:
    definition = (run.get("model") or {}).get("definition") or {}
    for category in definition.get("categories", []):
        indicators = [item for item in result["indicators"]
                      if any(d["key"] == item["key"] and d.get("category_key") == category["key"]
                             for d in definition.get("dimensions", []))]
        if not indicators:
            continue
        doc.add_heading(category["label"], 2)
        for indicator in indicators:
            if indicator["simulated"]:
                doc.add_paragraph(f"{indicator['label']}采用模拟数据，仅展示试算过程，不评价产品真实表现。")
                continue
            values = [(product, row) for product in PRODUCTS for row in real_scores(result, product)
                      if row["key"] == indicator["key"]]
            if not values:
                doc.add_paragraph(f"{indicator['label']}数据尚未齐备，待补充后评价。")
                continue
            maximum = max(row["ratio"] for _, row in values)
            minimum = min(row["ratio"] for _, row in values)
            top = "、".join(product for product, row in values if row["ratio"] == maximum)
            bottom = "、".join(product for product, row in values if row["ratio"] == minimum)
            doc.add_paragraph(f"{indicator['label']}：{len(values)} 条产品线已有可用评分。"
                              + ("各产品得分相同。" if maximum == minimum else f"当前得分最高为{top}，最低为{bottom}。")
                              + "评分反映本批次相对表现，具体原值、增幅和名次见同维度明细表。")


def product_radar(scores: list[dict]) -> BytesIO | None:
    if len(scores) < 3:
        return None
    font_file = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
    if not font_file.exists():
        return None
    font = font_manager.FontProperties(fname=str(font_file), size=8)
    angles = [2 * math.pi * i / len(scores) for i in range(len(scores))]
    values = [row["ratio"] for row in scores]
    fig, ax = plt.subplots(figsize=(4.9, 2.8), subplot_kw={"projection": "polar"})
    ax.set_theta_offset(math.pi / 2)
    ax.set_theta_direction(-1)
    ax.plot(angles + [angles[0]], values + [values[0]], color="#466654", linewidth=1)
    ax.fill(angles + [angles[0]], values + [values[0]], color="#466654", alpha=.12)
    ax.set_xticks(angles)
    ax.set_xticklabels([str(index) for index in range(1, len(scores) + 1)], fontproperties=font)
    ax.set_ylim(0, max([1, *values]) * 1.05)
    ax.set_yticks([.25, .5, .75, 1], ["25%", "50%", "75%", "100%"], fontsize=6, color="#707070")
    ax.grid(color="#d9d9d9", linewidth=.5)
    ax.spines["polar"].set_color("#d9d9d9")
    fig.subplots_adjust(left=.19, right=.81, top=.84, bottom=.16)
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=180)
    plt.close(fig)
    buffer.seek(0)
    return buffer


def add_product_evaluations(doc, run: dict, result: dict) -> None:
    draft = (run.get("notes") or {}).get("evaluation_draft") or {}
    evaluations = {row["product"]: row for row in draft.get("products", [])}
    summaries = {row["product"]: row for row in result.get("model_summary", {}).get("rows", [])}
    products = sorted(PRODUCTS, key=lambda product: (summaries.get(product, {}).get("rank") or 999, PRODUCTS.index(product)))
    doc.add_heading("附录 重点产品详细评价和提升建议", 1)
    doc.add_paragraph("本附录逐款汇总评分表现、雷达图、详细评价与提升建议。"
                      + ("评价文字由 AI 辅助起草，可由用户编辑，需由负责人审阅确认。" if evaluations else "评价依据已计算的真实数据；经营原因及处置方向由负责人补充确认。"))
    add_dimension_analysis(doc, run, result)
    section = doc.add_heading("分产品总体评价和提升建议", 2)
    section.paragraph_format.page_break_before = True
    for product_index, product in enumerate(products):
        summary = summaries.get(product, {})
        suffix = f" 总分{summary['total']:.1f} 排名{summary['rank']}" if result.get("model_summary", {}).get("ready") and summary.get("total") is not None and summary.get("rank") is not None else ""
        heading = doc.add_heading(product + suffix, 3)
        heading.paragraph_format.page_break_before = product_index > 0
        if result.get("result_status") == "trial" and suffix:
            doc.add_paragraph("以上为含模拟数据的试算总分与排名；下文仅依据真实数据评价。")
        scores = real_scores(result, product)
        chart = product_radar(scores)
        if scores:
            paragraph = doc.add_paragraph("评分表现：" + "；".join(f"{row['label']} {row['weighted_score']:.1f}/{row['weight']:g}分" for row in scores) + "。")
            paragraph.paragraph_format.keep_with_next = chart is not None
        if chart is not None:
            doc.add_picture(chart, width=Inches(4.5))
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            caption = doc.add_paragraph("雷达图按各指标得分占对应权重的比例展示，模拟和缺失指标不参与。")
            caption.paragraph_format.keep_with_next = True
            doc.add_paragraph("图中指标：" + "；".join(f"{index}={row['label']}" for index, row in enumerate(scores, 1)) + "。")
        evaluation = evaluations.get(product)
        if evaluation:
            doc.add_paragraph("详细评价：" + narrative_text(evaluation["evaluation"]))
            doc.add_paragraph("提升建议：" + narrative_text(evaluation["recommendation"]))
        elif scores:
            ordered = sorted(scores, key=lambda row: row["ratio"], reverse=True)
            if ordered[0]["ratio"] != ordered[-1]["ratio"]:
                doc.add_paragraph(f"详细评价：{ordered[0]['label']}是本产品当前得分率较高的指标；{ordered[-1]['label']}得分率相对较低，建议结合原值和增幅核查改善空间。不同权重按得分率比较。")
                doc.add_paragraph(f"提升建议：优先围绕{ordered[-1]['label']}核实影响因素，形成改善措施，并以该指标对应原值、增幅及得分跟踪成效。具体经营原因和资源安排待负责人确认。")
            else:
                doc.add_paragraph("详细评价：已计算指标的得分率相同，尚不能据此区分本产品的主要短板。")
                doc.add_paragraph("提升建议：结合业务数据补充经营原因，明确需改善的指标及措施。")
        else:
            doc.add_paragraph("详细评价及提升建议：真实数据尚不足，待补充后评价。")
    if evaluations:
        doc.add_paragraph(f"AI 草稿来源：{draft.get('provider', '')} {draft.get('model', '')}；生成时间 {draft.get('generated_at', '')}。报告按生成时保存的审阅内容导出。")
