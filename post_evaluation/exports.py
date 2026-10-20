from __future__ import annotations

from copy import copy
from io import BytesIO
import os
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from .store import data_root
os.environ.setdefault("MPLCONFIGDIR", str(data_root() / "matplotlib"))
(data_root() / "matplotlib").mkdir(parents=True, exist_ok=True)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import ECONOMIC_FIELDS, METRICS, PRODUCTS, ROLES, field_label
from .result_state import describe_result
from .score_presentation import step_format, step_label
from .report_sections import add_product_evaluations, add_dimension_analysis, add_insights


def _make_generated_workbook_legacy(run: dict, result: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    year = int(run["year"])
    traces = {(item["product"], item["metric"]): item for item in result.get("traces", [])}
    weights = _metric_weights(run)

    _add_dimension_index(wb, year, weights)
    _add_market_sheet(wb, year, traces, weights["market_share"])
    _add_benefit_sheet(wb, year, traces, weights)
    _add_three_eff_sheet(wb, year, traces, weights["three_eff"])
    _add_self_research_sheet(wb, year, traces, weights)
    _add_quality_sheet(wb, year, traces, weights["quality"])
    _add_total_sheet(wb, year, result)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.print_title_rows = "1:2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


_GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
_THIN_BLACK = Side(style="thin", color="000000")
_MEDIUM_BLACK = Side(style="medium", color="000000")
_GRID_BORDER = Border(left=_THIN_BLACK, right=_THIN_BLACK, top=_THIN_BLACK, bottom=_THIN_BLACK)
_HEADER_BORDER = Border(left=_MEDIUM_BLACK, right=_MEDIUM_BLACK, top=_MEDIUM_BLACK, bottom=_MEDIUM_BLACK)
_BODY_FONT = Font(name="Arial Unicode MS", size=10, color="000000")
_HEADER_FONT = Font(name="Arial Unicode MS", size=10, bold=True, color="000000")
_TITLE_FONT = Font(name="Arial Unicode MS", size=14, bold=True, color="000000")
_PENDING = "待完成"


def _metric_weights(run: dict) -> dict[str, float]:
    weights = {key: 10.0 for key, _, _ in METRICS}
    definition = ((run.get("model") or {}).get("definition") or {})
    for indicator in definition.get("dimensions") or []:
        metric_keys = indicator.get("metric_keys") or []
        if not metric_keys:
            continue
        per_metric = float(indicator.get("weight", 0)) / len(metric_keys)
        for key in metric_keys:
            if key in weights:
                weights[key] = per_metric
    return weights


def _trace_step(trace: dict | None, index: int) -> dict:
    if not trace:
        return {}
    steps = trace.get("steps") or []
    return steps[index] if index < len(steps) else {}


def _ranked_products(traces: dict, metric: str) -> list[str]:
    order = {product: index for index, product in enumerate(PRODUCTS)}
    return sorted(PRODUCTS, key=lambda product: (
        (traces.get((product, metric)) or {}).get("rank") or 10_000,
        order[product],
    ))


def _score_value(value):
    return _PENDING if value is None else value


def _format_weight(value: float):
    return int(value) if float(value).is_integer() else value


def _setup_template_sheet(ws, title: str, columns: int, header_height: float = 54) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = _TITLE_FONT
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 27
    ws.row_dimensions[2].height = header_height
    ws.freeze_panes = "C3"


def _style_data_table(ws, first_row: int, last_row: int, columns: int, green_columns: set[int] | None = None) -> None:
    green_columns = green_columns or set()
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=columns):
        for cell in row:
            cell.font = _HEADER_FONT if cell.row == first_row else _BODY_FONT
            cell.border = _HEADER_BORDER if cell.row == first_row else _GRID_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.column in green_columns:
                cell.fill = _GREEN_FILL
    for row in ws.iter_rows(min_row=first_row + 1, max_row=last_row, min_col=1, max_col=columns):
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = Font(name="Arial Unicode MS", size=10, color="FF0000")


def _set_widths(ws, widths: dict[int, float], default: float = 12) -> None:
    for column in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = widths.get(column, default)


def _append_weight(ws, row: int, weight: float, note: str | None = None) -> None:
    ws.cell(row, 1, "权重")
    ws.cell(row, 2, _format_weight(weight))
    for cell in ws[row][:2]:
        cell.font = _BODY_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if note:
        ws.cell(row + 1, 2, note)
        ws.cell(row + 1, 2).font = Font(name="Arial Unicode MS", size=9, color="666666")


def _add_dimension_index(wb, year: int, weights: dict[str, float]) -> None:
    ws = wb.create_sheet("重点产品-雷达图维度")
    _setup_template_sheet(ws, "重点产品评估指标", 6, 28)
    ws.append(["序号", "维度", "一级指标", "二级指标", "来源", "权重"])
    short_year = str(year)[-2:]
    rows = [
        (1, "市场份额", "市场份额", f"{short_year}年市场份额", "咨询机构报告", weights["market_share"]),
        (2, "效益贡献", "效益贡献", f"{short_year}年收入", "财务部分产品核算", weights["revenue"]),
        (3, "效益贡献", "效益贡献", f"{str(year - 1)[-2:]}-{short_year}年收入增幅", "财务部分产品核算", None),
        (4, "效益贡献", "效益贡献", f"{short_year}年毛利", "财务部分产品核算", weights["gross_profit"]),
        (5, "效益贡献", "效益贡献", f"{str(year - 1)[-2:]}-{short_year}年毛利增幅", "财务部分产品核算", None),
        (6, "效益贡献", "效益贡献", f"{short_year}年毛利率", "财务部分产品核算", weights["gross_margin"]),
        (7, "效益贡献", "效益贡献", f"{str(year - 1)[-2:]}-{short_year}年毛利率增幅", "财务部分产品核算", None),
        (8, "效益贡献", "效益贡献", f"{short_year}年营业利润额", "财务部分产品核算", weights["operating_profit"]),
        (9, "效益贡献", "效益贡献", f"{str(year - 1)[-2:]}-{short_year}年营业利润增幅", "财务部分产品核算", None),
        (10, "效益贡献", "效益贡献", f"{short_year}年营业利润率", "财务部分产品核算", weights["operating_margin"]),
        (11, "效益贡献", "效益贡献", f"{str(year - 1)[-2:]}-{short_year}年营业利润率增幅", "财务部分产品核算", None),
        (12, "产品三效", "产品三效", "劳产", "财务部分产品核算", weights["three_eff"]),
        (13, "产品三效", "产品三效", "元效", "财务部分产品核算", None),
        (14, "产品三效", "产品三效", "人效", "财务部分产品核算", None),
        (15, "自研能力", "自研能力", "自研收入占比", "财务部分产品核算", weights["self_research"]),
        (16, "自研能力", "自研能力", "研发投入产出比", "财务部分产品核算", None),
        (17, "自研能力", "能力评测情况", "能力评测情况", "实验室评测结果", weights["ability"]),
        (18, "品质服务", "品质服务", "产品推广满意度", f"品质部{short_year}年推广产品满意度", weights["quality"]),
    ]
    for row in rows:
        ws.append(list(row))
    category_weights = {
        "市场份额": weights["market_share"],
        "效益贡献": sum(weights[key] for key in ("revenue", "gross_profit", "gross_margin", "operating_profit", "operating_margin")),
        "产品三效": weights["three_eff"],
        "自研能力": weights["self_research"] + weights["ability"],
        "品质服务": weights["quality"],
    }
    merge_specs = [
        (3, 3, "市场份额"),
        (4, 13, "效益贡献"),
        (14, 16, "产品三效"),
        (17, 19, "自研能力"),
        (20, 20, "品质服务"),
    ]
    for start, end, label in merge_specs:
        ws.cell(start, 2, f"{label}（{_format_weight(category_weights[label])}分）")
        if start != end:
            ws.merge_cells(start_row=start, start_column=2, end_row=end, end_column=2)
    for start, end in ((4, 13), (14, 16), (17, 18)):
        ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
    for start, end in ((4, 5), (6, 7), (8, 9), (10, 11), (12, 13), (14, 16), (17, 18)):
        ws.merge_cells(start_row=start, start_column=6, end_row=end, end_column=6)
    ws.append(["总分", None, None, None, None, _format_weight(sum(weights.values()))])
    _style_data_table(ws, 2, ws.max_row, 6)
    for cell in ws[ws.max_row]:
        cell.font = _HEADER_FONT
    _set_widths(ws, {1: 8, 2: 21, 3: 20, 4: 28, 5: 29, 6: 10})


def _add_market_sheet(wb, year: int, traces: dict, weight: float) -> None:
    ws = wb.create_sheet("重点产品-市场份额")
    _setup_template_sheet(ws, "重点产品-市场份额", 9, 58)
    short_previous = str(year - 1)[-2:]
    short_year = str(year)[-2:]
    ws.append([
        "序号", "产品线", f"市场份额\n（{year - 2}年）", f"市场份额\n（{year - 1}年）",
        f"市场份额\n（{year}年）", f"{short_previous}-{short_year}增长率", f"{short_year}年市场份额排名",
        f"得分\n（依据{short_year}年市场份额排名得分）", "备注",
    ])
    for number, product in enumerate(_ranked_products(traces, "market_share"), 1):
        trace = traces.get((product, "market_share"))
        step = _trace_step(trace, 0)
        ws.append([number, product, None, None, _score_value(step.get("value")), None,
                   _score_value((trace or {}).get("rank")), _score_value((trace or {}).get("score")), None])
    last_data_row = ws.max_row
    _style_data_table(ws, 2, last_data_row, 9, {8})
    for row in range(3, last_data_row + 1):
        for column in (3, 4, 5, 6):
            ws.cell(row, column).number_format = "0.00%"
        ws.cell(row, 8).number_format = "0.0"
    _append_weight(ws, last_data_row + 2, weight, "历史市场份额列按附件模板保留；当前评分口径只接入评估年度市场份额。")
    _set_widths(ws, {1: 8, 2: 18, 3: 16, 4: 16, 5: 16, 6: 15, 7: 16, 8: 23, 9: 28})


def _economic_group_headers(label: str, year: int, include_year_minus_two: bool) -> list[str]:
    short_previous = str(year - 1)[-2:]
    short_year = str(year)[-2:]
    value_label = "营业利润额" if label == "营业利润" else label
    headers = [f"{short_previous}年{value_label}", f"{short_year}年{label}"]
    if include_year_minus_two:
        headers.insert(0, f"{str(year - 2)[-2:]}年{value_label}")
    headers.extend([
        f"{short_year}年{value_label}排名", "得分", f"{short_previous}-{short_year}年{label}增幅",
        f"{short_previous}-{short_year}年{label}增幅排名", "得分", f"{label}/增幅平均分",
        f"{label}/增幅平均分排名", "得分",
    ])
    return headers


def _economic_group_values(trace: dict | None, include_year_minus_two: bool) -> list:
    previous = _trace_step(trace, 0)
    current = _trace_step(trace, 1)
    growth = _trace_step(trace, 2)
    average = _trace_step(trace, 3)
    values = [_score_value(previous.get("value")), _score_value(current.get("value"))]
    if include_year_minus_two:
        values.insert(0, None)
    values.extend([
        _score_value(current.get("rank")), _score_value(current.get("score")),
        _score_value(growth.get("value")), _score_value(growth.get("rank")), _score_value(growth.get("score")),
        _score_value(average.get("value")), _score_value((trace or {}).get("rank")), _score_value((trace or {}).get("score")),
    ])
    return values


def _add_benefit_sheet(wb, year: int, traces: dict, weights: dict[str, float]) -> None:
    ws = wb.create_sheet("重点产品-效益贡献")
    groups = [
        ("revenue", "收入", True),
        ("gross_profit", "毛利", True),
        ("gross_margin", "毛利率", True),
        ("operating_profit", "营业利润", False),
        ("operating_margin", "营业利润率", False),
    ]
    headers = ["序号", "产品线"]
    for _, label, include_year_minus_two in groups:
        headers.extend(_economic_group_headers(label, year, include_year_minus_two))
    _setup_template_sheet(ws, "重点产品-效益贡献", len(headers), 70)
    ws.append(headers)
    for number, product in enumerate(_ranked_products(traces, "revenue"), 1):
        values = [number, product]
        for metric, _, include_year_minus_two in groups:
            values.extend(_economic_group_values(traces.get((product, metric)), include_year_minus_two))
        ws.append(values)
    last_data_row = ws.max_row
    green_columns: set[int] = set()
    column = 3
    for _, _, include_year_minus_two in groups:
        size = 11 if include_year_minus_two else 10
        green_columns.update({column + size - 2, column + size - 1})
        column += size
    _style_data_table(ws, 2, last_data_row, len(headers), green_columns)
    column = 3
    for metric, _, include_year_minus_two in groups:
        size = 11 if include_year_minus_two else 10
        value_columns = [column + (1 if include_year_minus_two else 0), column + (2 if include_year_minus_two else 1)]
        growth_column = column + (5 if include_year_minus_two else 4)
        average_column = column + size - 3
        value_format = "0.00%" if metric in {"gross_margin", "operating_margin"} else "0.00"
        for row in range(3, last_data_row + 1):
            for value_column in value_columns:
                ws.cell(row, value_column).number_format = value_format
            ws.cell(row, growth_column).number_format = "0.00%"
            for score_column in (column + (4 if include_year_minus_two else 3), column + (7 if include_year_minus_two else 6), average_column, column + size - 1):
                ws.cell(row, score_column).number_format = "0.0"
        column += size
    note_row = last_data_row + 1
    ws.cell(note_row, 1, f"* {year - 2}年原值列按附件模板保留；当前系统评分输入为{year - 1}年和{year}年，未接入内容留空。")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
    ws.cell(note_row, 1).font = Font(name="Arial Unicode MS", size=9, color="666666")
    _append_weight(ws, note_row + 1, sum(weights[key] for key, _, _ in groups))
    widths = {1: 8, 2: 18}
    for col in range(3, len(headers) + 1):
        widths[col] = 14
    _set_widths(ws, widths)


def _add_three_eff_sheet(wb, year: int, traces: dict, weight: float) -> None:
    ws = wb.create_sheet("重点产品-产品三效")
    headers = ["序号", "产品线", "劳产", "排名", "得分", "元效", "排名", "得分", "人效", "排名", "得分", "三效平均分", "排名", "得分"]
    _setup_template_sheet(ws, "重点产品-产品三效", len(headers), 48)
    ws.append(headers)
    for number, product in enumerate(_ranked_products(traces, "three_eff"), 1):
        trace = traces.get((product, "three_eff"))
        steps = [_trace_step(trace, index) for index in range(4)]
        ws.append([number, product,
                   _score_value(steps[0].get("value")), _score_value(steps[0].get("rank")), _score_value(steps[0].get("score")),
                   _score_value(steps[1].get("value")), _score_value(steps[1].get("rank")), _score_value(steps[1].get("score")),
                   _score_value(steps[2].get("value")), _score_value(steps[2].get("rank")), _score_value(steps[2].get("score")),
                   _score_value(steps[3].get("value")), _score_value((trace or {}).get("rank")), _score_value((trace or {}).get("score"))])
    last_data_row = ws.max_row
    _style_data_table(ws, 2, last_data_row, len(headers), {13, 14})
    for row in range(3, last_data_row + 1):
        for column in (3, 5, 6, 8, 9, 11, 12, 14):
            ws.cell(row, column).number_format = "0.00" if column in (3, 6, 9) else "0.0"
    _append_weight(ws, last_data_row + 2, weight)
    _set_widths(ws, {1: 8, 2: 18, 12: 17}, 12)


def _add_self_research_sheet(wb, year: int, traces: dict, weights: dict[str, float]) -> None:
    ws = wb.create_sheet("重点产品-自研能力")
    headers = [
        "序号", "产品线", "自研收入占比", "排名", "得分", "研发总成本", "收入", "研发投入产出比", "排名", "得分",
        "自研占比/ROI平均分", "自研占比/ROI平均分排名", "得分",
        f"{str(year - 2)[-2:]}年能力评测项", "得分", f"{str(year - 1)[-2:]}年对标评测项", "得分",
        f"{str(year - 1)[-2:]}年能力评测项", "得分", f"{str(year)[-2:]}年能力评测项", "得分",
        f"{str(year)[-2:]}年AI对标评测项", "等级", "得分", "能力评测平均分", "能力评测排名", "得分",
    ]
    _setup_template_sheet(ws, "重点产品-自研能力", len(headers), 70)
    ws.append(headers)
    for number, product in enumerate(_ranked_products(traces, "self_research"), 1):
        self_trace = traces.get((product, "self_research"))
        ability_trace = traces.get((product, "ability"))
        ratio = _trace_step(self_trace, 0)
        revenue = _trace_step(self_trace, 1)
        total_cost = _trace_step(self_trace, 9)
        roi = _trace_step(self_trace, 10)
        average = _trace_step(self_trace, 11)
        ability = _trace_step(ability_trace, 0)
        ability_label = "系统模拟占位值" if ability_trace else _PENDING
        ws.append([
            number, product,
            _score_value(ratio.get("value")), _score_value(ratio.get("rank")), _score_value(ratio.get("score")),
            _score_value(total_cost.get("value")), _score_value(revenue.get("value")), _score_value(roi.get("value")),
            _score_value(roi.get("rank")), _score_value(roi.get("score")), _score_value(average.get("value")),
            _score_value((self_trace or {}).get("rank")), _score_value((self_trace or {}).get("score")),
            None, None, None, None, None, None,
            ability_label, _score_value(ability.get("value")), None, None, None,
            _score_value(ability.get("value")), _score_value((ability_trace or {}).get("rank")), _score_value((ability_trace or {}).get("score")),
        ])
    last_data_row = ws.max_row
    _style_data_table(ws, 2, last_data_row, len(headers), {12, 13, 26, 27})
    for row in range(3, last_data_row + 1):
        ws.cell(row, 3).number_format = "0.00%"
        for column in (5, 6, 7, 8, 10, 11, 13, 21, 25, 27):
            ws.cell(row, column).number_format = "0.00" if column in (6, 7, 8) else "0.0"
    note_row = last_data_row + 1
    ws.cell(note_row, 1, "* 自研收入占比与研发投入产出比来自本批次上传数据；历年能力评测列按附件模板保留，当前能力评测为系统模拟占位值。")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
    ws.cell(note_row, 1).font = Font(name="Arial Unicode MS", size=9, color="666666")
    _append_weight(ws, note_row + 1, weights["self_research"] + weights["ability"])
    widths = {1: 8, 2: 18, 3: 17, 6: 16, 7: 16, 8: 18, 11: 19, 12: 20}
    for col in range(14, 28):
        widths[col] = 18 if col not in (15, 17, 19, 21, 23, 24, 26, 27) else 10
    _set_widths(ws, widths, 11)


def _add_quality_sheet(wb, year: int, traces: dict, weight: float) -> None:
    ws = wb.create_sheet("重点产品-品质服务")
    headers = ["序号", "产品线", "产品推广满意度", "排名", "得分"]
    _setup_template_sheet(ws, "重点产品-品质服务", len(headers), 32)
    ws.append(headers)
    for number, product in enumerate(_ranked_products(traces, "quality"), 1):
        trace = traces.get((product, "quality"))
        step = _trace_step(trace, 0)
        ws.append([number, product, _score_value(step.get("value")), _score_value((trace or {}).get("rank")), _score_value((trace or {}).get("score"))])
    last_data_row = ws.max_row
    _style_data_table(ws, 2, last_data_row, len(headers), {5})
    for row in range(3, last_data_row + 1):
        ws.cell(row, 3).number_format = "0.0"
        ws.cell(row, 5).number_format = "0.0"
    _append_weight(ws, last_data_row + 2, weight, f"未接入{year}年产品推广满意度时保留“待完成”，不按 0 分处理。")
    _set_widths(ws, {1: 8, 2: 20, 3: 24, 4: 10, 5: 14})


def _add_total_sheet(wb, year: int, result: dict) -> None:
    ws = wb.create_sheet("重点产品-总分")
    headers = [
        "序号", "产品线", f"{str(year)[-2:]}年市场份额", "收入/增幅", "毛利/增幅", "毛利率/增幅",
        "营业利润/增幅", "营业利润率/增幅", "产品三效", "自研占比/ROI", "能力评测", "品质服务", "总分", "排名",
    ]
    _setup_template_sheet(ws, "重点产品-总分", len(headers), 38)
    ws.append(headers)
    summary = result.get("model_summary") or {}
    summary_by_product = {item["product"]: item for item in summary.get("rows", [])}
    product_order = {product: index for index, product in enumerate(PRODUCTS)}
    products = sorted(PRODUCTS, key=lambda product: (
        (summary_by_product.get(product) or {}).get("rank") or 10_000,
        product_order[product],
    ))
    metric_keys = [key for key, _, _ in METRICS]
    for number, product in enumerate(products, 1):
        product_scores = result.get("scores", {}).get(product, {})
        product_summary = summary_by_product.get(product) or {}
        ws.append([
            number, product,
            *[_score_value(product_scores.get(key)) for key in metric_keys],
            _score_value(product_summary.get("total") if product_summary.get("complete") else None),
            _score_value(product_summary.get("rank") if product_summary.get("complete") else None),
        ])
    last_data_row = ws.max_row
    _style_data_table(ws, 2, last_data_row, len(headers), {13, 14})
    for row in range(3, last_data_row + 1):
        for column in range(3, 14):
            ws.cell(row, column).number_format = "0.0"
        ws.cell(row, 14).number_format = "0"
    _set_widths(ws, {1: 8, 2: 18, 3: 18, 4: 17, 5: 14, 6: 17, 7: 18, 8: 19, 9: 15, 10: 18, 11: 15, 12: 14, 13: 12, 14: 10})


_EXPORT_TEMPLATE = Path(__file__).resolve().parent / "templates" / "product-evaluation-export-template.xlsx"
_TEMPLATE_FIRST_DATA_ROW = 3
_TEMPLATE_LAST_DATA_ROW = 15


def _write_template_rows(ws, rows: list[dict[int, object]]) -> None:
    """Write values into the template's fixed 13-row body without touching formatting."""
    capacity = _TEMPLATE_LAST_DATA_ROW - _TEMPLATE_FIRST_DATA_ROW + 1
    if len(rows) > capacity:
        raise ValueError(f"导出模板最多支持 {capacity} 条产品线，当前为 {len(rows)} 条。")
    for row_number in range(_TEMPLATE_FIRST_DATA_ROW, _TEMPLATE_LAST_DATA_ROW + 1):
        ws.cell(row_number, 1, row_number - _TEMPLATE_FIRST_DATA_ROW + 1)
        for column in range(2, ws.max_column + 1):
            ws.cell(row_number, column, None)
    for offset, values in enumerate(rows):
        row_number = _TEMPLATE_FIRST_DATA_ROW + offset
        for column, value in values.items():
            ws.cell(row_number, column, value)


def _template_ranked_products(traces: dict, metric: str) -> list[str]:
    product_order = {product: index for index, product in enumerate(PRODUCTS)}

    def key(product: str) -> tuple[float, int]:
        rank = (traces.get((product, metric)) or {}).get("rank")
        return (float(rank) if isinstance(rank, (int, float)) else float("inf"), product_order[product])

    return sorted(PRODUCTS, key=key)


def _economic_template_values(trace: dict | None, start_column: int, include_year_minus_two: bool) -> dict[int, object]:
    previous = _trace_step(trace, 0)
    current = _trace_step(trace, 1)
    growth = _trace_step(trace, 2)
    average = _trace_step(trace, 3)
    column = start_column + (1 if include_year_minus_two else 0)
    return {
        column: _score_value(previous.get("value")),
        column + 1: _score_value(current.get("value")),
        column + 2: _score_value(current.get("rank")),
        column + 3: _score_value(current.get("score")),
        column + 4: _score_value(growth.get("value")),
        column + 5: _score_value(growth.get("rank")),
        column + 6: _score_value(growth.get("score")),
        column + 7: _score_value(average.get("value")),
        column + 8: _score_value((trace or {}).get("rank")),
        column + 9: _score_value((trace or {}).get("score")),
    }


def _fill_market_template(wb, traces: dict) -> None:
    rows = []
    for product in _template_ranked_products(traces, "market_share"):
        trace = traces.get((product, "market_share"))
        current = _trace_step(trace, 0)
        rows.append({
            2: product,
            5: _score_value(current.get("value")),
            7: _score_value((trace or {}).get("rank")),
            8: _score_value((trace or {}).get("score")),
        })
    _write_template_rows(wb["重点产品-市场份额"], rows)


def _fill_benefit_template(wb, traces: dict) -> None:
    groups = (
        ("revenue", 3, True),
        ("gross_profit", 14, True),
        ("gross_margin", 25, True),
        ("operating_profit", 36, False),
        ("operating_margin", 46, False),
    )
    rows = []
    for product in _template_ranked_products(traces, "revenue"):
        values: dict[int, object] = {2: product}
        for metric, start_column, include_year_minus_two in groups:
            values.update(_economic_template_values(traces.get((product, metric)), start_column, include_year_minus_two))
        rows.append(values)
    _write_template_rows(wb["重点产品-效益贡献"], rows)


def _fill_three_eff_template(wb, traces: dict) -> None:
    rows = []
    for product in _template_ranked_products(traces, "three_eff"):
        trace = traces.get((product, "three_eff"))
        steps = [_trace_step(trace, index) for index in range(4)]
        rows.append({
            2: product,
            3: _score_value(steps[0].get("value")), 4: _score_value(steps[0].get("rank")), 5: _score_value(steps[0].get("score")),
            6: _score_value(steps[1].get("value")), 7: _score_value(steps[1].get("rank")), 8: _score_value(steps[1].get("score")),
            9: _score_value(steps[2].get("value")), 10: _score_value(steps[2].get("rank")), 11: _score_value(steps[2].get("score")),
            12: _score_value(steps[3].get("value")), 13: _score_value((trace or {}).get("rank")), 14: _score_value((trace or {}).get("score")),
        })
    _write_template_rows(wb["重点产品-产品三效"], rows)


def _fill_self_research_template(wb, traces: dict) -> None:
    rows = []
    for product in _template_ranked_products(traces, "self_research"):
        self_trace = traces.get((product, "self_research"))
        ability_trace = traces.get((product, "ability"))
        ratio = _trace_step(self_trace, 0)
        revenue = _trace_step(self_trace, 1)
        total_cost = _trace_step(self_trace, 9)
        roi = _trace_step(self_trace, 10)
        average = _trace_step(self_trace, 11)
        ability = _trace_step(ability_trace, 0)
        rows.append({
            2: product,
            3: _score_value(ratio.get("value")), 4: _score_value(ratio.get("rank")), 5: _score_value(ratio.get("score")),
            6: _score_value(total_cost.get("value")), 7: _score_value(revenue.get("value")),
            8: _score_value(roi.get("value")), 9: _score_value(roi.get("rank")), 10: _score_value(roi.get("score")),
            11: _score_value(average.get("value")), 12: _score_value((self_trace or {}).get("rank")), 13: _score_value((self_trace or {}).get("score")),
            20: "系统模拟占位值" if ability_trace else _PENDING,
            21: _score_value(ability.get("value")), 25: _score_value(ability.get("value")),
            26: _score_value((ability_trace or {}).get("rank")), 27: _score_value((ability_trace or {}).get("score")),
        })
    _write_template_rows(wb["重点产品-自研能力"], rows)


def _fill_quality_template(wb, traces: dict) -> None:
    rows = []
    for product in _template_ranked_products(traces, "quality"):
        trace = traces.get((product, "quality"))
        current = _trace_step(trace, 0)
        rows.append({
            2: product,
            3: _score_value(current.get("value")),
            4: _score_value((trace or {}).get("rank")),
            5: _score_value((trace or {}).get("score")),
        })
    _write_template_rows(wb["重点产品-品质服务"], rows)


def _fill_total_template(wb, result: dict) -> None:
    summary_by_product = {
        item["product"]: item for item in (result.get("model_summary") or {}).get("rows", [])
    }
    product_order = {product: index for index, product in enumerate(PRODUCTS)}
    products = sorted(PRODUCTS, key=lambda product: (
        (summary_by_product.get(product) or {}).get("rank")
        if isinstance((summary_by_product.get(product) or {}).get("rank"), (int, float))
        else float("inf"),
        product_order[product],
    ))
    rows = []
    metric_keys = [key for key, _, _ in METRICS]
    for product in products:
        scores = result.get("scores", {}).get(product, {})
        summary = summary_by_product.get(product) or {}
        values: dict[int, object] = {2: product}
        for column, metric in enumerate(metric_keys, 3):
            values[column] = _score_value(scores.get(metric))
        values[13] = _score_value(summary.get("total") if summary.get("complete") else None)
        values[14] = _score_value(summary.get("rank") if summary.get("complete") else None)
        rows.append(values)
    _write_template_rows(wb["重点产品-总分"], rows)


def _copy_cell_style(target, source) -> None:
    # Style IDs belong to a workbook-local table, so copy the components
    # rather than assigning ``_style`` across workbooks.
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _safe_sheet_title(label: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", f"重点产品-{label}").strip()[:31] or "重点产品-维度"
    title = base
    suffix = 2
    while title in existing:
        addition = f"-{suffix}"
        title = f"{base[:31-len(addition)]}{addition}"
        suffix += 1
    existing.add(title)
    return title


def _export_style_sources(template):
    detail = template["重点产品-产品三效"]
    total = template["重点产品-总分"]
    return {
        "title": detail["A1"],
        "header": detail["B2"],
        "score_header": detail["N2"],
        "body": detail["C3"],
        "text_body": detail["B3"],
        "score_body": detail["N3"],
        "total_body": total["M3"],
    }


def _prepare_dynamic_sheet(ws, title: str, headers: list[str], score_columns: set[int], styles: dict) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 2))
    ws.cell(1, 1, title)
    _copy_cell_style(ws.cell(1, 1), styles["title"])
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 58
    for column, label in enumerate(headers, 1):
        cell = ws.cell(2, column, label)
        _copy_cell_style(cell, styles["score_header"] if column in score_columns else styles["header"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"
    ws.sheet_view.showGridLines = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:2"


def _export_step_label(indicator: dict, metric: str, step_index: int, step: dict, data_year: int) -> str:
    return step_label(indicator, metric, step_index, step, data_year)


def _trace_column_specs(indicator: dict, traces: dict, result: dict, data_year: int) -> list[dict]:
    specs: list[dict] = []
    for metric in indicator.get("metric_keys") or []:
        type_id = indicator.get("calculation_type") or metric
        sample = next((traces.get((product, metric)) for product in PRODUCTS if traces.get((product, metric))), None)
        if not sample:
            label = f"{indicator.get('label')}得分"
            specs.append({"label": label, "kind": "score", "metric": metric, "role": "score"})
            continue
        for step_index, step in enumerate(sample.get("steps") or []):
            # 自研能力只展示汇总后的研发总成本，成本组成仍保留在计算追溯数据中。
            if type_id == "self_research" and 2 <= step_index <= 8:
                continue
            label = _export_step_label(indicator, type_id, step_index, step, data_year)
            specs.append({"label": label, "kind": "step", "metric": metric, "calculation_type": type_id, "step": step_index, "field": "value", "role": "value", "unit": step_format(type_id, step_index, step)["unit"]})
            if any((_trace_step(traces.get((product, metric)), step_index)).get("rank") is not None for product in PRODUCTS):
                specs.append({"label": f"{label}排名", "kind": "step", "metric": metric, "step": step_index, "field": "rank", "role": "rank"})
            if any((_trace_step(traces.get((product, metric)), step_index)).get("score") is not None for product in PRODUCTS):
                specs.append({"label": f"{label}得分", "kind": "step", "metric": metric, "step": step_index, "field": "score", "role": "score"})
    return specs


def _column_value(spec: dict, product: str, traces: dict, result: dict):
    if spec["kind"] == "score":
        return (result.get("scores", {}).get(product) or {}).get(spec["metric"])
    return _trace_step(traces.get((product, spec["metric"])), spec["step"]).get(spec["field"])


def _style_dynamic_body(ws, row: int, column: int, role: str, styles: dict, label: str, value_format: str | None = None, unit: str = "") -> None:
    source = styles["score_body"] if role == "score" else styles["text_body"] if column == 2 else styles["body"]
    _copy_cell_style(ws.cell(row, column), source)
    ws.cell(row, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    value_format = value_format or _result_column_format(role, label)
    if value_format == "integer":
        ws.cell(row, column).number_format = "0"
    elif value_format == "score":
        ws.cell(row, column).number_format = "0.0"
    elif value_format == "percent":
        ws.cell(row, column).number_format = "0.00%"
    elif value_format in {"number", "ratio"}:
        ws.cell(row, column).number_format = "0.00"
    if value_format in {"number", "ratio"} and (unit or value_format == "ratio"):
        suffix = (unit or "倍").replace(chr(34), "")
        ws.cell(row, column).number_format += f'" {suffix}"'


def _indicator_rank(product: str, indicator: dict, traces: dict) -> float:
    for metric in indicator.get("metric_keys") or []:
        rank = (traces.get((product, metric)) or {}).get("rank")
        if isinstance(rank, (int, float)):
            return float(rank)
    return float("inf")


def _result_column_format(role: str, label: str) -> str:
    if role in {"index", "rank"}:
        return "integer"
    if role == "score":
        return "score"
    if role == "text":
        return "text"
    if any(word in label for word in ("份额", "占比", "毛利率", "利润率", "增幅")):
        return "percent"
    return "number"


def _spec_result_format(spec: dict) -> str:
    if spec["role"] in {"rank", "score"}:
        return _result_column_format(spec["role"], spec["label"])
    return step_format(spec.get("calculation_type") or spec.get("metric"), spec.get("step", -1), {})["format"]


def make_result_tables(run: dict, result: dict) -> list[dict]:
    """Build the canonical table model shared by the score page and Excel export."""
    definition = ((run.get("model") or {}).get("definition") or {})
    indicators = [item for item in definition.get("dimensions") or [] if item.get("enabled", True)]
    categories = [
        category for category in definition.get("categories") or []
        if any(item.get("category_key") == category.get("key") for item in indicators)
    ]
    traces = {(item["product"], item["metric"]): item for item in result.get("traces", [])}
    evaluation_year = int(run["year"])
    data_year = int(run.get("data_year") if run.get("data_year") is not None else evaluation_year)
    product_order = {product: index for index, product in enumerate(PRODUCTS)}
    tables: list[dict] = []

    summary_by_product = {
        item["product"]: item for item in (result.get("model_summary") or {}).get("rows", [])
    }
    total_products = sorted(PRODUCTS, key=lambda product: (
        (summary_by_product.get(product) or {}).get("rank")
        if isinstance((summary_by_product.get(product) or {}).get("rank"), (int, float))
        else float("inf"),
        product_order[product],
    ))
    total_columns = [
        {"label": "序号", "role": "index", "format": "integer", "trace_metric": None},
        {"label": "产品线", "role": "text", "format": "text", "trace_metric": None},
        *[
            {
                "label": item.get("label") or item.get("key"),
                "role": "score",
                "format": "score",
                "trace_metric": (item.get("metric_keys") or [None])[0],
            }
            for item in indicators
        ],
        {"label": "总分", "role": "score", "format": "score", "trace_metric": None},
        {"label": "排名", "role": "rank", "format": "integer", "trace_metric": None},
    ]
    total_rows = []
    for number, product in enumerate(total_products, 1):
        summary = summary_by_product.get(product) or {}
        dimension_scores = {item.get("key"): item.get("score") for item in summary.get("dimensions", [])}
        total_rows.append({
            "product": product,
            "values": [
                number,
                product,
                *[dimension_scores.get(item.get("key")) for item in indicators],
                summary.get("total") if summary.get("complete") else None,
                summary.get("rank") if summary.get("complete") else None,
            ],
        })
    tables.append({
        "key": "total",
        "label": "总分",
        "title": f"{evaluation_year}年重点产品后评估 · 总分",
        "kind": "total",
        "columns": total_columns,
        "rows": total_rows,
    })

    for category in categories:
        grouped = [item for item in indicators if item.get("category_key") == category.get("key")]
        specs = [
            spec
            for indicator in grouped
            for spec in _trace_column_specs(indicator, traces, result, data_year)
        ]
        sort_key = category.get("export_sort_indicator_key")
        sort_indicator = next((item for item in grouped if item.get("key") == sort_key), grouped[0])
        products = sorted(
            PRODUCTS,
            key=lambda product: (_indicator_rank(product, sort_indicator, traces), product_order[product]),
        )
        columns = [
            {"label": "序号", "role": "index", "format": "integer", "trace_metric": None},
            {"label": "产品线", "role": "text", "format": "text", "trace_metric": None},
            *[
                {
                    "label": spec["label"],
                    "role": spec["role"],
                    "format": _spec_result_format(spec),
                    "unit": spec.get("unit", ""),
                    "trace_metric": spec.get("metric"),
                }
                for spec in specs
            ],
        ]
        rows = [
            {
                "product": product,
                "values": [
                    number,
                    product,
                    *[_column_value(spec, product, traces, result) for spec in specs],
                ],
            }
            for number, product in enumerate(products, 1)
        ]
        tables.append({
            "key": str(category.get("key") or category.get("label")),
            "label": str(category.get("label") or "维度"),
            "title": f"重点产品-{category.get('label')}",
            "kind": "dimension",
            "columns": columns,
            "rows": rows,
        })
    return tables


def _add_dynamic_index(wb, year: int, categories: list[dict], indicators: list[dict], styles: dict) -> None:
    ws = wb.create_sheet("重点产品-雷达图维度")
    headers = ["序号", "维度", "指标", "权重", "计算规则", "导出排序指标"]
    _prepare_dynamic_sheet(ws, f"{year}年重点产品后评估 · 模型维度", headers, set(), styles)
    indicator_by_key = {item.get("key"): item for item in indicators}
    row = 3
    for category in categories:
        grouped = [item for item in indicators if item.get("category_key") == category.get("key")]
        sort_key = category.get("export_sort_indicator_key")
        if sort_key not in {item.get("key") for item in grouped}:
            sort_key = grouped[0].get("key") if grouped else None
        for indicator in grouped:
            values = [row - 2, category.get("label"), indicator.get("label"), _format_weight(float(indicator.get("weight", 0))), indicator.get("rule_summary") or "—", indicator_by_key.get(sort_key, {}).get("label") or "—"]
            for column, value in enumerate(values, 1):
                ws.cell(row, column, value)
                _style_dynamic_body(ws, row, column, "value", styles, headers[column - 1])
            row += 1
    for column, width in enumerate((8, 18, 22, 10, 64, 20), 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def _add_dynamic_dimension_sheet(wb, table: dict, styles: dict, existing: set[str]) -> None:
    headers = [column["label"] for column in table["columns"]]
    score_columns = {index for index, column in enumerate(table["columns"], 1) if column["role"] == "score"}
    title = _safe_sheet_title(table["label"], existing)
    ws = wb.create_sheet(title)
    _prepare_dynamic_sheet(ws, table["title"], headers, score_columns, styles)
    for number, item in enumerate(table["rows"], 1):
        row = number + 2
        values = item["values"]
        for column, value in enumerate(values, 1):
            ws.cell(row, column, _score_value(value))
            role = table["columns"][column - 1]["role"]
            _style_dynamic_body(ws, row, column, role, styles, headers[column - 1], table["columns"][column - 1]["format"], table["columns"][column - 1].get("unit", ""))
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    for column, header in enumerate(headers[2:], 3):
        ws.column_dimensions[get_column_letter(column)].width = min(28, max(13, len(str(header)) * 1.8))


def _add_dynamic_total(wb, table: dict, styles: dict) -> None:
    ws = wb.create_sheet("重点产品-总分")
    headers = [column["label"] for column in table["columns"]]
    score_columns = {index for index, column in enumerate(table["columns"], 1) if column["role"] == "score"}
    _prepare_dynamic_sheet(ws, table["title"], headers, score_columns, styles)
    for number, item in enumerate(table["rows"], 1):
        row = number + 2
        values = item["values"]
        for column, value in enumerate(values, 1):
            ws.cell(row, column, _score_value(value))
            role = table["columns"][column - 1]["role"]
            _style_dynamic_body(ws, row, column, role, styles, headers[column - 1], table["columns"][column - 1]["format"], table["columns"][column - 1].get("unit", ""))
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    for column in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 17


def make_workbook(run: dict, result: dict) -> bytes:
    """Generate one sheet per current model dimension using template styling."""
    result = describe_result((run.get("model") or {}).get("definition"), result)
    if not _EXPORT_TEMPLATE.exists():
        raise FileNotFoundError(f"导出模板不存在：{_EXPORT_TEMPLATE}")
    template = load_workbook(_EXPORT_TEMPLATE, data_only=False, keep_links=False)
    styles = _export_style_sources(template)
    wb = Workbook()
    wb.remove(wb.active)
    definition = ((run.get("model") or {}).get("definition") or {})
    indicators = [item for item in definition.get("dimensions") or [] if item.get("enabled", True)]
    categories = [
        category for category in definition.get("categories") or []
        if any(item.get("category_key") == category.get("key") for item in indicators)
    ]
    evaluation_year = int(run["year"])
    _add_dynamic_index(wb, evaluation_year, categories, indicators, styles)
    index_sheet = wb["重点产品-雷达图维度"]
    note_row = index_sheet.max_row + 2
    for label, value in (("雷达图口径", "各指标加权贡献÷该指标满分；0–100%表示满分占比，零权重不绘制。"),
                         ("结果状态", result["status_note"]),
                         ("指标覆盖", f"{result['calculated_indicator_count']} / {result['indicator_count']} 项；权重覆盖 {result['covered_points']:g} / {result['weight_total']:g} 分"),
                         ("缺失指标", "、".join(result["missing_indicators"]) or "无"),
                         ("模拟指标", "、".join(result["simulated_indicators"]) or "无")):
        index_sheet.cell(note_row, 1, label)
        index_sheet.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=6)
        index_sheet.cell(note_row, 2, value).alignment = Alignment(wrap_text=True, vertical="center")
        index_sheet.row_dimensions[note_row].height = 32
        note_row += 1
    tables = make_result_tables(run, result)
    existing = set(wb.sheetnames) | {"重点产品-总分"}
    for table in tables[1:]:
        _add_dynamic_dimension_sheet(wb, table, styles, existing)
    _add_dynamic_total(wb, tables[0], styles)
    total_sheet = wb["重点产品-总分"]
    total_sheet.cell(total_sheet.max_row + 2, 1, result["status_note"])
    template.close()
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def make_report(run: dict, result: dict) -> bytes:
    definition = (run.get("model") or {}).get("definition")
    result = describe_result(definition, result)
    available = [(item["key"], item["label"]) for item in result["indicators"] if item["complete"]]
    pending = result["missing_indicators"]
    summary = result.get("model_summary") or {}
    summary_by_product = {item["product"]: item for item in summary.get("rows", [])}
    totals_ready = bool(summary.get("ready"))
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.78)
    sec.bottom_margin = Inches(.7)
    sec.left_margin = Inches(.7)
    sec.right_margin = Inches(.7)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial Unicode MS")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 3"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    title_border = styles["Title"]._element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        styles["Title"]._element.pPr.remove(title_border)
    doc.add_paragraph(f"{run['year']}年重点产品后评估报告", style="Title")
    doc.add_paragraph(("分项测算稿" if result["result_status"] == "partial" else "试算稿（含模拟数据）" if result["result_status"] == "trial" else "完整测算稿") + " · 待负责人审阅")
    coverage_note = result["status_note"]
    doc.add_paragraph(f"评估范围：13 条重点产品线。规则版本 {run.get('rule_version', 'v1-2025')}。当前已计算 {len(available)}/{result['indicator_count']} 项指标，覆盖 {result['covered_points']:g}/{result['weight_total']:g} 分。{coverage_note}")
    draft = (run.get("notes") or {}).get("evaluation_draft") or {}
    composed = draft.get("version") == 2 and bool(draft.get("overview"))
    if composed:
        doc.add_heading("一 总体结论", 1)
        add_insights(doc, draft["overview"])
        doc.add_heading("二 分维度分析", 1)
        add_dimension_analysis(doc, run, result)
        doc.add_page_break()
        add_product_evaluations(doc, run, result, main_section=True)
        doc.add_page_break()
        doc.add_heading("附录 评估口径与评分结果", 1)
    else:
        doc.add_heading("一 上年度评估改善情况", 1)
        doc.add_paragraph("本批次尚未录入上年度评估排名、改善措施及完成情况。该部分留待负责人补充依据并审阅，不根据本年度分数推断改善效果。")
        doc.add_heading("二 本年度产品后评估结果", 1)
        doc.add_heading("总体评价", 2)
        doc.add_paragraph(coverage_note)
        ranked = sorted((row for row in summary.get("rows", []) if row.get("rank") is not None), key=lambda row: row["rank"])
        if totals_ready and ranked:
            doc.add_paragraph("本批次综合得分前列产品为" + "、".join(row["product"] for row in ranked[:3])
                              + "。评分仅在本批次重点产品内比较；各产品业务表现、短板和提升建议见详细评价附录。")
        else:
            doc.add_paragraph("当前先报告已计算指标的表现，待数据齐备后形成综合评价。各产品已具备的数据及待补充事项见详细评价附录。")
    doc.add_heading("评估口径与数据来源", 2)
    for indicator in (definition or {}).get("dimensions", []):
        if indicator.get("enabled", True):
            doc.add_paragraph(f"{indicator.get('label')}（权重{indicator.get('weight'):g}分）：{indicator.get('rule_summary')}")
    doc.add_paragraph("采用竞争排名，原始赋分按指标权重折算为贡献分。缺失数据不以0分补齐；完整计算步骤与来源见配套评分总表。")
    for file in run["files"]:
        doc.add_paragraph(f"• {file['name']}｜{ROLES.get(file['role'], file['role'])}｜{file['sha256'][:12]}…")
    doc.add_page_break()
    doc.add_heading("重点产品评分结果", 2)
    table = doc.add_table(rows=1, cols=1 + len(available) + (2 if totals_ready else 0))
    table.style = "Light Shading Accent 1"
    table.autofit = False
    numeric_columns = len(available) + (2 if totals_ready else 0)
    score_width = 5.9 / max(numeric_columns, 1)
    widths = [Inches(1.15)] + [Inches(score_width)] * numeric_columns
    for column, width in zip(table.columns, widths):
        column.width = width
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    headers = ["产品线"] + [label for _, label in available] + (["总分", "排名"] if totals_ready else [])
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for product in PRODUCTS:
        row = table.add_row().cells
        for cell, width in zip(row, widths):
            cell.width = width
        row[0].text = product
        for index, (key, _) in enumerate(available, 1):
            score = next((item.get("score") for item in summary_by_product.get(product, {}).get("dimensions", []) if item.get("key") == key), None)
            row[index].text = f"{score:.1f}" if score is not None else "—"
        if totals_ready:
            product_summary = summary_by_product[product]
            row[len(available) + 1].text = f"{product_summary['total']:.1f}"
            row[len(available) + 2].text = str(product_summary["rank"])
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for text_run in paragraph.runs:
                    text_run.font.name = "Arial Unicode MS"
                    text_run.font.size = Pt(8)
                    text_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial Unicode MS")
    doc.add_paragraph("注：" + coverage_note + ("" if totals_ready else "不生成正式总分。"))
    if available:
        fig, ax = plt.subplots(figsize=(7.1, 2.6))
        labels = [f"{i+1}" for i in range(len(available))]
        averages = [sum(next((item.get("score", 0) for item in summary_by_product.get(p, {}).get("dimensions", []) if item.get("key") == key), 0) for p in PRODUCTS) / len(PRODUCTS) for key, _ in available]
        bars = ax.bar(labels, averages, color="#6d8b75", width=.55)
        ax.bar_label(bars, fmt="%.1f", padding=3, fontsize=8)
        ax.set_ylim(0, max([float(item["weight"]) for item in result["indicators"] if item["complete"]] + [1]) * 1.06)
        font_file = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
        if font_file.exists():
            chinese_font = font_manager.FontProperties(fname=str(font_file))
            ax.set_ylabel("平均得分", fontsize=9, fontproperties=chinese_font)
            ax.set_title("已计算维度平均得分", loc="left", fontsize=10, fontproperties=chinese_font)
        else:
            ax.set_ylabel("Average score", fontsize=9)
            ax.set_title("Average by calculated dimension", loc="left", fontsize=10)
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", color="#e4e9e8", linewidth=.5)
        ax.set_axisbelow(True)
        fig.tight_layout()
        chart = BytesIO()
        fig.savefig(chart, format="png", dpi=180)
        plt.close(fig)
        chart.seek(0)
        doc.add_picture(chart, width=Inches(6.7))
        doc.add_paragraph("图注：" + "；".join(f"{i+1}={label}" for i, (_, label) in enumerate(available)) + "。展示加权贡献分。" + coverage_note)
    doc.add_heading("数据校验与待确认事项", 2)
    errors = [issue for issue in result["issues"] if issue["level"] == "error"]
    warnings = [issue for issue in result["issues"] if issue["level"] == "warning"]
    doc.add_paragraph(f"数据错误 {len(errors)} 项，复核提示 {len(warnings)} 项。")
    for issue in errors[:10]:
        doc.add_paragraph(f"• {issue['message']}")
    if len(errors) > 10:
        doc.add_paragraph("其余错误见评分总表的“校验与待确认”工作表。")
    warning_groups = {}
    for issue in warnings:
        warning_groups.setdefault((issue.get("code"), issue.get("field")), []).append(issue)
    for (_, field), group in warning_groups.items():
        products = list(dict.fromkeys(issue.get("product") for issue in group if issue.get("product")))
        doc.add_paragraph(f"• {field_label(field) if field else '数据复核'}：{'、'.join(products)}。"
                          + f"提示示例：{group[0]['message']}。完整提示见配套评分总表。")
    doc.add_paragraph(f"待完成指标：{'、'.join(pending) or '无'}。缺少上年排名、产品矩阵与人员调整依据，本稿不自动编造对应结论。")
    notes = run.get("notes") or {}
    supplements = [(heading, (notes.get(key) or "").strip()) for heading, key in (
        ("经营原因分析", "analysis"), ("产品处置建议", "recommendations"), ("报告备注", "review"),
    ) if (notes.get(key) or "").strip()]
    if supplements:
        doc.add_heading("附录 补充说明" if composed else "三 补充说明", 1)
        for heading, content in supplements:
            doc.add_heading(heading, 2)
            doc.add_paragraph(content)
    if not composed:
        doc.add_page_break()
        add_product_evaluations(doc, run, result)
    doc.add_heading("附录 计算明细与来源", 1)
    doc.add_paragraph("每条产品线的原值、增幅、两级名次、得分和来源单元格见配套评分总表的“计算明细”及“数据来源”工作表。")
    footer = sec.footer.paragraphs[0]
    footer.text = "产品后评估 · 阶段性测算稿 · 人工结论须审阅"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
