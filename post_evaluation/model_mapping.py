"""Match workbook headers against the current model's named raw inputs."""
from __future__ import annotations

from collections import defaultdict
import re
import unicodedata

from openpyxl import load_workbook

from .calculation_types import source_metrics
from .constants import field_label, input_specs_for_dimension
from .ingest import _header_years, auto_mapping, detect_metrics
from .input_names import bare_input_name, declared_input_names, normalize_name


def _header(value: str) -> tuple[str, set[int], str | None]:
    label = unicodedata.normalize("NFKC", str(value or ""))
    years = _header_years(label)
    period = "previous_year" if re.match(r"\s*(?:上一年度|上一年|上年|去年)", label) else None
    if re.match(r"\s*(?:当年|本年|评估年度|最新年度|本年度)", label):
        period = "current_year"
    label = re.sub(r"(?:19|20)\d{2}年?|(?<!\d)\d{2}年", "", label)
    label = bare_input_name(label).strip().strip("()[]【】_-/ ")
    return normalize_name(label), years, period


def match_model_workbook(path, dimensions: list[dict], year: int, targets: list[str] | None = None) -> tuple[list[str], list[str], dict]:
    """Use explicit model names, retaining known source-layout compatibility.

    Only validated engine input slots can be mapped; arbitrary header text
    never creates a new calculation. Duplicate columns/sheets stay unresolved.
    """
    active = [item for item in dimensions if item.get("enabled", True) and (targets is None or item["key"] in targets)]
    legacy_metrics = detect_metrics(path)
    # BytesIO is reusable after openpyxl has moved its cursor.
    mapping_metrics = list(dict.fromkeys(metric for item in active for metric in source_metrics(item))) if targets is not None else legacy_metrics
    legacy = auto_mapping(path, mapping_metrics, year)
    by_name = defaultdict(set)
    named_owners = defaultdict(set)
    custom = set()
    custom_targets = set()
    for dimension in active:
        names = declared_input_names(dimension, year)
        for spec in input_specs_for_dimension(dimension, year):
            key = spec["key"]
            name = names.get(key, field_label(key))
            by_name[normalize_name(bare_input_name(name))].add(key)
            named_owners[(normalize_name(bare_input_name(name)), key)].add(dimension["key"])
            if normalize_name(bare_input_name(name)) != normalize_name(bare_input_name(field_label(key))):
                custom.add(key)
                custom_targets.add(dimension["key"])
    workbook = load_workbook(path, read_only=True, data_only=False)
    candidates = []
    matched_targets = set()
    try:
        for sheet in workbook:
            context = _header_years(sheet.title)
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20)):
                product_cols = [cell.column_letter for cell in row if normalize_name(str(cell.value or "")) in {"产品线", "产品名称", "公司产品", "产品"}]
                if not product_cols:
                    context.update(y for cell in row for y in _header_years(str(cell.value or "")))
                    continue
                if len(product_cols) != 1:
                    continue
                matches = defaultdict(list)
                recognized = set()
                for cell in row:
                    label, years, period = _header(cell.value)
                    for key in by_name.get(label, set()):
                        recognized.add(key)
                        matched_targets.update(named_owners[(label, key)])
                        expected_year = int(key.split(":")[1])
                        relative_year = year - 1 if period == "previous_year" else year
                        if (years and (years != {expected_year} or (context and expected_year not in context))) or (not years and context and context != {expected_year}):
                            continue
                        if period and relative_year != expected_year:
                            continue
                        if not years and not context and not period and expected_year != year:
                            continue
                        matches[key].append(cell.column_letter)
                if recognized:
                    candidates.append({"sheet": sheet.title, "header_row": row[0].row,
                                       "ambiguous_fields": [key for key in recognized if len(matches[key]) > 1],
                                       "product_col": product_cols[0], "fields": {key: columns[0] if len(columns) == 1 else None for key in recognized for columns in [matches[key]]}})
    finally:
        workbook.close()
    named_fields = {key for candidate in candidates for key in candidate["fields"]}
    # Existing multirow finance/source layouts keep their mature reader unless
    # a model-specific business name actually needs the new matching path.
    use_named = bool(candidates) and (not any(legacy["fields"].values()) or bool(named_fields & custom))
    detected_targets = {item["key"] for item in active if item["key"] not in custom_targets and set(source_metrics(item)) & set(legacy_metrics)}
    if targets is not None:
        detected_targets.update(targets)
    if use_named:
        detected_targets.update(matched_targets)
        if len(candidates) == 1:
            mapping = candidates[0]
            if legacy.get("sheet") == mapping["sheet"] and legacy.get("header_row") == mapping["header_row"]:
                mapping["fields"] = {**legacy["fields"], **mapping["fields"]}
        else:
            mapping = {"sheet": None, "header_row": 1, "product_col": "A", "fields": {key: None for key in named_fields}}
            mapping["confirmation_note"] = "多个表头或工作表匹配所需数据，请确认工作表和数据列"
    else:
        mapping = legacy
    selected = [item for item in active if item["key"] in detected_targets]
    expected = {spec["key"] for item in selected for spec in input_specs_for_dimension(item, year)}
    # Preserve legacy display-only columns (e.g. per-file ratios) as before.
    for key in expected:
        mapping["fields"].setdefault(key, None)
    metrics = list(dict.fromkeys(metric for item in selected for metric in source_metrics(item)))
    if mapping.get("ambiguous_fields"):
        mapping["confirmation_note"] = "同一数据字段有多个候选列，请确认数据来源"
    return [item["key"] for item in selected], metrics, mapping
